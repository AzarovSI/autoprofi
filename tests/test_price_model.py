"""«Цены → Ценовая модель»: вводимые параметры с историей (price_model_hist). Только локальная wb_price_test."""
import datetime
import io
import unittest
from zoneinfo import ZoneInfo

import psycopg
from fastapi.testclient import TestClient

from test_wb_prices import prepare_local_db
from test_access_roles import reset_roles
from app import auth, db, permissions, price_model
from app.main import app

TODAY = datetime.datetime.now(ZoneInfo("Europe/Moscow")).date()
YESTERDAY = TODAY - datetime.timedelta(days=1)


def reset_price_model(category_rows=()):
    """Состояние рабочей БД до миграций модели: лист 13.10 (price_model_category); category_rows —
    строки price_category_hist, если нужно проверить перенос (тогда первая миграция уже применена)."""
    with psycopg.connect(db.DB_DSN, autocommit=True) as conn:
        for t in ("price_model_hist", "price_category_hist", "price_model_category", "price_model_rules"):
            conn.execute(f"DROP TABLE IF EXISTS {t}")
        if conn.execute("SELECT to_regclass('app_schema_migrations')").fetchone()[0]:
            conn.execute("DELETE FROM app_schema_migrations WHERE name = ANY(%s)", (price_model.MIGRATIONS,))
        conn.execute("""CREATE TABLE price_model_category (
            seller_article text NOT NULL, marketplace text NOT NULL, category_price integer NOT NULL,
            PRIMARY KEY (seller_article, marketplace))""")
        if category_rows:
            path = "migrations/2026_10_21_price_category.sql"
            conn.execute(open(path, encoding="utf-8").read())
            conn.execute("INSERT INTO app_schema_migrations(name) VALUES (%s) ON CONFLICT DO NOTHING",
                         (price_model.MIGRATIONS[0],))
            for row in category_rows:
                conn.execute("INSERT INTO price_category_hist (seller_article, marketplace, date, category_price) "
                             "VALUES (%s, %s, %s, %s)", row)


class PriceModelTest(unittest.TestCase):
    category_rows = ()

    def setUp(self):
        prepare_local_db()
        reset_roles()
        permissions.migrate()
        reset_price_model(self.category_rows)
        price_model.migrate()
        price_model.migrate()  # повторный запуск ничего не меняет
        ids = {r["name"]: r["id"] for r in db.query_all("SELECT id, name FROM app_roles")}
        # Роль «Только цена категории»: смотреть модель и менять цену категории, без ДРР, маржи и порогов.
        cat_role = db.query_one("INSERT INTO app_roles (name) VALUES ('Только цена категории') RETURNING id")["id"]
        db.execute("INSERT INTO app_role_permissions (role_id, permission) VALUES (%s, 'login'), "
                   "(%s, 'prices.model.view'), (%s, 'prices.model.edit')", (cat_role, cat_role, cat_role))
        db.execute("INSERT INTO app_users (id, email, display_name, role, role_id) VALUES "
                   "(101, 'admin@local.test', 'Админ', 'admin', %s), (102, 'rop@local.test', 'РОП', 'user', %s), "
                   "(103, 'cat@local.test', 'Продакт', 'user', %s)",
                   (ids["Администратор"], ids["РОП"], cat_role))
        self.client = TestClient(app, raise_server_exceptions=True)

    def tearDown(self):
        db.execute("DELETE FROM app_users WHERE id >= 100")
        reset_roles()

    def call(self, uid, method, path, **kw):
        token = auth.create_token({"id": uid, "email": ""})
        return self.client.request(method, path, headers={"Authorization": f"Bearer {token}"}, **kw)

    def set_(self, mp, value, param="category_price", art="sku-1", uid=101):
        return self.call(uid, "POST", "/api/prices/model_params",
                         json={"seller_article": art, "marketplace": mp, "param": param, "value": value})

    def items(self, date_to=None):
        path = "/api/prices/model_params" + (f"?date_to={date_to.isoformat()}" if date_to else "")
        r = self.call(101, "GET", path)
        self.assertEqual(r.status_code, 200, r.text)
        return [(i["seller_article"], i["marketplace"], i["param"], i["value"], i["updated_by"])
                for i in r.json()["items"]]

    def hist(self):
        return [(r["marketplace"], r["param"], r["date"], None if r["value"] is None else float(r["value"]))
                for r in db.query_all("SELECT marketplace, param, date, value FROM price_model_hist "
                                      "ORDER BY marketplace, param, date")]


class PriceModelParamsTest(PriceModelTest):
    def test_old_tables_dropped(self):
        for t in ("price_model_category", "price_category_hist"):
            self.assertIsNone(db.query_one("SELECT to_regclass(%s) AS t", (t,))["t"])

    def test_three_params_and_last_value_of_day(self):
        self.assertEqual(self.set_("wb", 970).status_code, 200)
        self.assertEqual(self.set_("wb", 0.125, "drr_pct").status_code, 200)
        self.assertEqual(self.set_("ozon", -0.05, "margin_pct").status_code, 200)
        # Повторная правка в тот же день перезаписывает запись дня, а не добавляет новую.
        r = self.set_("wb", 990)
        self.assertEqual((r.json()["value"], r.json()["date"]), (990, TODAY.isoformat()))
        self.assertEqual(self.hist(), [("Ozon", "margin_pct", TODAY, -0.05),
                                       ("Wildberries", "category_price", TODAY, 990.0),
                                       ("Wildberries", "drr_pct", TODAY, 0.125)])
        # Артикул канонизируется, цена — целым числом, проценты — долями, видно автора.
        self.assertEqual(self.items(), [("SKU-1", "Ozon", "margin_pct", -0.05, "Админ"),
                                        ("SKU-1", "Wildberries", "category_price", 990, "Админ"),
                                        ("SKU-1", "Wildberries", "drr_pct", 0.125, "Админ")])
        self.assertIsInstance(self.items()[1][3], int)

    def test_history_slice_on_period_end(self):
        db.execute("INSERT INTO price_model_hist (seller_article, marketplace, param, date, value) "
                   "VALUES ('SKU-1', 'Wildberries', 'drr_pct', %s, 0.1)", (YESTERDAY,))
        self.set_("wb", 0.15, "drr_pct")
        self.assertEqual(self.items(YESTERDAY), [("SKU-1", "Wildberries", "drr_pct", 0.1, None)])
        self.assertEqual(self.items(TODAY), [("SKU-1", "Wildberries", "drr_pct", 0.15, "Админ")])
        self.assertEqual(self.items(YESTERDAY - datetime.timedelta(days=1)), [])
        # Удаление — запись дня с пустым значением: сегодня значения нет, вчерашний срез не меняется.
        self.assertEqual(self.set_("wb", None, "drr_pct").status_code, 200)
        self.assertEqual(self.items(TODAY), [])
        self.assertEqual(self.items(YESTERDAY), [("SKU-1", "Wildberries", "drr_pct", 0.1, None)])
        self.assertEqual(self.hist(), [("Wildberries", "drr_pct", YESTERDAY, 0.1),
                                       ("Wildberries", "drr_pct", TODAY, None)])

    def test_unchanged_value_does_not_grow_history(self):
        db.execute("INSERT INTO price_model_hist (seller_article, marketplace, param, date, value) "
                   "VALUES ('SKU-1', 'Ozon', 'margin_pct', %s, 0.2)", (YESTERDAY,))
        self.set_("ozon", 0.2, "margin_pct")
        self.assertEqual(self.hist(), [("Ozon", "margin_pct", YESTERDAY, 0.2)])
        # Изменили и в тот же день вернули прежнее — запись дня убирается.
        self.set_("ozon", 0.25, "margin_pct")
        self.set_("ozon", 0.2, "margin_pct")
        self.assertEqual(self.hist(), [("Ozon", "margin_pct", YESTERDAY, 0.2)])
        # Удаление несуществующего значения ничего не пишет.
        self.set_("ozon", None, "drr_pct")
        self.assertEqual(self.hist(), [("Ozon", "margin_pct", YESTERDAY, 0.2)])

    def test_rejects_bad_values(self):
        bad = (("wb", 0, "category_price"), ("wb", -5, "category_price"), ("wb", 10.5, "category_price"),
               ("foo", 100, "category_price"), ("wb", 100, "price"),
               ("wb", -0.001, "drr_pct"), ("wb", 1.001, "drr_pct"), ("wb", 0.1234, "drr_pct"),
               ("ozon", -0.501, "margin_pct"), ("ozon", 1.1, "margin_pct"))
        for mp, value, param in bad:
            with self.subTest(mp=mp, value=value, param=param):
                r = self.set_(mp, value, param)
                self.assertEqual(r.status_code, 400, r.text)
        self.assertEqual(self.items(), [])
        # Не число (Infinity / NaN в JSON) — 400, а не ошибка сервера.
        for raw in ('Infinity', 'NaN'):
            with self.subTest(raw=raw):
                r = self.client.request(
                    "POST", "/api/prices/model_params",
                    content='{"seller_article": "SKU-1", "marketplace": "wb", "param": "drr_pct", "value": %s}' % raw,
                    headers={"Authorization": "Bearer " + auth.create_token({"id": 101, "email": ""}),
                             "Content-Type": "application/json"})
                self.assertEqual(r.status_code, 400, r.text)
        # Границы допустимы.
        for mp, value, param in (("wb", 0, "drr_pct"), ("wb", 1, "margin_pct"), ("ozon", -0.5, "margin_pct")):
            with self.subTest(mp=mp, value=value, param=param):
                self.assertEqual(self.set_(mp, value, param).status_code, 200)

    def test_each_param_has_own_right(self):
        # У роли только право на цену категории: цену меняет, ДРР и маржу — нет, пороги — нет.
        self.assertEqual(self.set_("wb", 1000, uid=103).status_code, 200)
        self.assertEqual(self.set_("wb", 0.1, "drr_pct", uid=103).status_code, 403)
        self.assertEqual(self.set_("wb", 0.2, "margin_pct", uid=103).status_code, 403)
        r = self.call(103, "PUT", "/api/prices/model_rules", json={"rules": {"Ozon": 0.02}})
        self.assertEqual(r.status_code, 403, r.text)
        self.assertEqual(self.items(), [("SKU-1", "Wildberries", "category_price", 1000, "Продакт")])

    def test_rules_default_save_and_validate(self):
        r = self.call(101, "GET", "/api/prices/model_params")
        self.assertEqual(r.json()["rules"], {"Wildberries": 0.05, "Ozon": 0.05})
        self.assertEqual(r.json()["rules_defaults"], {"Wildberries": 0.05, "Ozon": 0.05})
        r = self.call(101, "PUT", "/api/prices/model_rules", json={"rules": {"Ozon": 0.025}})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["rules"], {"Wildberries": 0.05, "Ozon": 0.025})
        for bad in ({"Ozon": 0}, {"Ozon": 0.6}, {"Ozon": 0.0125}, {"Yandex": 0.05}, {"Ozon": "x"}):
            with self.subTest(bad=bad):
                r = self.call(101, "PUT", "/api/prices/model_rules", json={"rules": bad})
                self.assertEqual(r.status_code, 400, r.text)
        self.assertEqual(self.call(101, "GET", "/api/prices/model_params").json()["rules"],
                         {"Wildberries": 0.05, "Ozon": 0.025})

    def test_spp_on_category_price_date(self):
        # Подсказка сигнала: СПП из ценового отчёта за день, с которого действует цена категории.
        with db.transaction() as tx:
            tx.execute("""INSERT INTO mp_price_daily (seller_article, marketplace, date, upload_price, buyer_price,
                              spp_pct, spp_source_date, upload_source, spp_is_estimated)
                          VALUES ('SKU-1', 'Wildberries', %s, 1400, 980, .3, %s, 'upload_wb_spp', false),
                                 ('SKU-1', 'Ozon', %s, 1800, 990, .45, %s, 'upload_ozon_coinvest', true)""",
                       (TODAY, TODAY, TODAY, TODAY))
        self.set_("wb", 1000)
        self.set_("ozon", 1100)
        self.set_("wb", 0.1, "drr_pct")
        got = {(i["marketplace"], i["param"]): i["spp_at_date"]
               for i in self.call(101, "GET", "/api/prices/model_params").json()["items"]}
        # WB — СПП за этот день; Ozon — оценочная не берётся; у ДРР СПП не нужна.
        self.assertEqual(got, {("Wildberries", "category_price"): 0.3, ("Ozon", "category_price"): None,
                               ("Wildberries", "drr_pct"): None})

    def test_only_with_rights(self):
        # У РОП нет функций «Ценовой модели» — ни смотреть, ни менять.
        self.assertEqual(self.call(102, "GET", "/api/prices/model_params").status_code, 403)
        r = self.call(102, "POST", "/api/prices/model_params",
                      json={"seller_article": "SKU-1", "marketplace": "wb", "param": "drr_pct", "value": 0.1})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.items(), [])


def xlsx(rows, header=None, formats=None):
    """Файл шаблона: шапка как у выгрузки, строки — списки значений; formats — {(строка, колонка): формат}."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header or (list(price_model.EXCEL_REF) + [c[0] for c in price_model.excel_columns()]))
    for r in rows:
        ws.append(r)
    for (ri, ci), fmt in (formats or {}).items():
        ws.cell(ri, ci).number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(art, drr_wb=None, margin_wb=None, cat_wb=None, drr_oz=None, margin_oz=None, cat_oz=None):
    return [None, None, None, art, None, drr_wb, margin_wb, cat_wb, drr_oz, margin_oz, cat_oz]


class PriceModelExcelTest(PriceModelTest):
    """Шаблон Excel и его загрузка (этап 2 плана владельца 21.10.2026)."""

    def setUp(self):
        super().setUp()
        # SKU-1 — на WB и Ozon, SKU-2 — только на WB (prepare_local_db).
        db.execute("INSERT INTO catalog_marketplace (seller_article, marketplace) VALUES ('SKU-1', 'Ozon')")

    def export(self, articles=None, uid=101):
        body = {"date_from": "2026-09-01", "date_to": TODAY.isoformat()}
        if articles is not None:
            body["articles"] = articles
        return self.call(uid, "POST", "/api/prices/model_export", json=body)

    def upload(self, data, uid=101):
        return self.call(uid, "POST", "/api/prices/model_import",
                         files={"file": ("m.xlsx", data, "application/octet-stream")})

    def test_export_template(self):
        self.set_("wb", 1000)
        self.set_("wb", 0.125, "drr_pct")
        self.set_("ozon", -0.05, "margin_pct")
        import openpyxl
        r = self.export(["SKU-1"])
        self.assertEqual(r.status_code, 200, r.text)
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), list(price_model.EXCEL_REF) + [
            "ДРР WB", "Маржа WB", "Цена категории WB", "ДРР Ozon", "Маржа Ozon", "Цена категории Ozon"])
        # Только переданные страницей артикулы; значения — срез на конец периода.
        self.assertEqual(rows[1:], [("Продукция ECOM", "Тестовая группа", "Товары", "SKU-1", "SKU-1",
                                     0.125, None, 1000, None, -0.05, None)])
        self.assertEqual(ws.cell(2, 6).number_format, "0.0%")
        self.assertEqual(ws.cell(2, 8).number_format, "#,##0")
        self.assertEqual(ws.cell(2, 4).data_type, "s")
        # Без списка — все товары отчёта за период.
        all_rows = list(openpyxl.load_workbook(io.BytesIO(self.export().content)).active.iter_rows(values_only=True))
        self.assertEqual([x[3] for x in all_rows[1:]], ["SKU-1", "SKU-2"])

    def test_import_formats_catalog_and_counts(self):
        data = xlsx([
            # Процентный формат — уже доля; число без формата и текст — проценты.
            row("sku-1", 0.125, 20, "1 500", "8,5%", None, 1100),
            # SKU-2 есть только на WB: значение Ozon пропускается, WB — пишется.
            row("SKU-2", cat_wb=900, cat_oz=950),
            # Неизвестный артикул не заводится.
            row("NEW-1", cat_wb=700),
        ], formats={(2, 6): "0.0%"})
        r = self.upload(data)
        self.assertEqual(r.status_code, 200, r.text)
        res = r.json()
        self.assertEqual((res["articles"], res["set"], res["cleared"]), (2, 6, 0))
        self.assertEqual((res["skipped_unknown"], res["unknown_sample"]), (1, ["NEW-1"]))
        self.assertEqual((res["skipped_mp"], res["skipped_mp_sample"]), (1, ["SKU-2 (Ozon)"]))
        self.assertEqual(res["skipped_columns"], [])
        self.assertEqual(self.items(), [
            ("SKU-1", "Ozon", "category_price", 1100, "Админ"), ("SKU-1", "Ozon", "drr_pct", 0.085, "Админ"),
            ("SKU-1", "Wildberries", "category_price", 1500, "Админ"), ("SKU-1", "Wildberries", "drr_pct", 0.125, "Админ"),
            ("SKU-1", "Wildberries", "margin_pct", 0.2, "Админ"), ("SKU-2", "Wildberries", "category_price", 900, "Админ")])
        self.assertIsNone(db.query_one("SELECT 1 AS x FROM catalog_marketplace WHERE seller_article = 'NEW-1'"))
        # Тот же файл ещё раз — ничего не меняется, история не растёт.
        n = db.query_one("SELECT count(*) AS n FROM price_model_hist")["n"]
        res = self.upload(data).json()
        self.assertEqual((res["set"], res["cleared"]), (0, 0))
        self.assertEqual(db.query_one("SELECT count(*) AS n FROM price_model_hist")["n"], n)

    def test_import_empty_cell_deletes(self):
        self.set_("wb", 1000)
        self.set_("wb", 0.1, "drr_pct")
        res = self.upload(xlsx([row("SKU-1", drr_wb=0.1, cat_wb=None)], formats={(2, 6): "0.0%"})).json()
        self.assertEqual((res["set"], res["cleared"]), (0, 1))
        self.assertEqual(self.items(), [("SKU-1", "Wildberries", "drr_pct", 0.1, "Админ")])

    def test_import_error_rejects_whole_file(self):
        self.set_("wb", 1000)
        data = xlsx([row("SKU-1", cat_wb=1200, drr_oz=1.5), row("SKU-2", drr_wb=150, cat_wb="abc"),
                     row("sku-1", cat_wb=1300)], formats={(2, 9): "0.0%"})
        r = self.upload(data)
        self.assertEqual(r.status_code, 422, r.text)
        detail = r.json()["detail"]
        # Значение — как его видит сотрудник: ячейка в процентах — «150%», а не «1.5».
        self.assertIn("строка 2, «ДРР Ozon»: «150%»", detail)
        self.assertIn("строка 3, «ДРР WB»: «150»", detail)
        self.assertIn("строка 3, «Цена категории WB»: «abc»", detail)
        self.assertIn("строка 4: артикул SKU-1 уже был в строке 2", detail)
        # Ничего не записано — даже верная строка.
        self.assertEqual(self.items(), [("SKU-1", "Wildberries", "category_price", 1000, "Админ")])

    def test_import_reads_only_allowed_columns(self):
        db.execute("""INSERT INTO app_role_permissions (role_id, permission)
                      SELECT role_id, 'prices.model.import' FROM app_users WHERE id = 103""")
        res = self.upload(xlsx([row("SKU-1", drr_wb=10, margin_wb=20, cat_wb=1000)]), uid=103).json()
        self.assertEqual(res["skipped_columns"], ["ДРР WB", "Маржа WB", "ДРР Ozon", "Маржа Ozon"])
        self.assertEqual(self.items(), [("SKU-1", "Wildberries", "category_price", 1000, "Продакт")])

    def test_import_bad_files(self):
        self.assertEqual(self.upload(b"not excel").status_code, 422)
        r = self.upload(xlsx([["SKU-1", 1000]], header=["Код", "Цена категории WB"]))
        self.assertIn("нет колонки «Артикул»", r.json()["detail"])
        r = self.upload(xlsx([["SKU-1", 1]], header=["Артикул", "Что-то"]))
        self.assertIn("нет колонок, которые вы можете загружать", r.json()["detail"])

    def test_excel_rights(self):
        self.assertEqual(self.export(uid=102).status_code, 403)
        self.assertEqual(self.upload(xlsx([]), uid=102).status_code, 403)


class SppHistoryTest(PriceModelTest):
    """История СПП по дням для подсказки у СПП в «Индексе цен» и «Ценовой модели» (22.10.2026)."""

    def test_days_back_from_date(self):
        with db.transaction() as tx:
            tx.execute("""INSERT INTO mp_price_daily (seller_article, marketplace, date, upload_price, buyer_price,
                              spp_pct, spp_source_date, upload_source, spp_is_estimated) VALUES
                          ('SKU-1', 'Wildberries', '2026-10-20', 1000, 700, .3, '2026-10-20', 'upload_wb_spp', false),
                          ('SKU-1', 'Wildberries', '2026-10-18', 1000, NULL, NULL, NULL, 'upload_wb_spp', false),
                          ('SKU-1', 'Wildberries', '2026-10-17', 1000, 720, .28, '2026-10-16', 'legacy', true),
                          ('SKU-1', 'Wildberries', '2026-10-05', 1000, 800, .2, '2026-10-05', 'upload_wb_spp', false),
                          ('SKU-1', 'Ozon', '2026-10-20', 1000, 500, .5, '2026-10-20', 'upload_ozon_coinvest', false)""")
        r = self.call(101, "GET", "/api/prices/spp_history?seller_article=SKU-1&marketplace=wb&date_to=2026-10-20&days=4")
        self.assertEqual(r.status_code, 200, r.text)
        # Свежие первыми; нет отчёта и пустая СПП — None; оценочная помечена; другая площадка и дни вне окна не попадают.
        self.assertEqual(r.json()["days"], [
            {"date": "2026-10-20", "spp": 0.3, "est": False}, {"date": "2026-10-19", "spp": None, "est": False},
            {"date": "2026-10-18", "spp": None, "est": False}, {"date": "2026-10-17", "spp": 0.28, "est": True}])
        self.assertEqual(len(self.call(101, "GET", "/api/prices/spp_history?seller_article=SKU-1&marketplace=ozon"
                                                   "&date_to=2026-10-20").json()["days"]), 14)

    def test_bad_params_and_rights(self):
        base = "/api/prices/spp_history?seller_article=SKU-1"
        self.assertEqual(self.call(101, "GET", base + "&marketplace=ya&date_to=2026-10-20").status_code, 400)
        self.assertEqual(self.call(101, "GET", base + "&marketplace=wb&date_to=").status_code, 400)
        self.assertEqual(self.call(101, "GET", base + "&marketplace=wb&date_to=2026-10-20&days=61").status_code, 422)
        # Продакт с правом «Ценовая модель» (без «Индекса цен») видит историю.
        self.assertEqual(self.call(103, "GET", base + "&marketplace=wb&date_to=2026-10-20").status_code, 200)


class PriceModelMoveTest(PriceModelTest):
    """Перенос цен категории, введённых после выкладки 21.10, из price_category_hist."""
    category_rows = (("SKU-1", "Wildberries", YESTERDAY, 900), ("SKU-1", "Wildberries", TODAY, None),
                     ("SKU-2", "Ozon", TODAY, 1500))

    def test_category_rows_moved(self):
        self.assertIsNone(db.query_one("SELECT to_regclass('price_category_hist') AS t")["t"])
        self.assertEqual(self.hist(), [("Ozon", "category_price", TODAY, 1500.0),
                                       ("Wildberries", "category_price", YESTERDAY, 900.0),
                                       ("Wildberries", "category_price", TODAY, None)])
        self.assertEqual(self.items(), [("SKU-2", "Ozon", "category_price", 1500, None)])


if __name__ == "__main__":
    unittest.main()
