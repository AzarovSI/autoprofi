# -*- coding: utf-8 -*-
"""Прайс-лист: дата начала действия базовой цены в Excel (с 12.10.2026, вариант А владельца)."""
import asyncio
import datetime as dt
import io
import unittest
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import HTTPException, UploadFile

from test_wb_prices import prepare_local_db
from app import db, daily_pi
from app.routers import prices_router as pr

MSK = ZoneInfo("Europe/Moscow")
ADMIN = {"id": 1, "role": "admin", "is_admin": True}


def ts(y, m, d):
    return dt.datetime(y, m, d, tzinfo=MSK)


def ver(price, start, end=None):
    return {"base_price": price, "valid_from": start, "valid_to": end, "updated_by_user_id": 1}


class ApplyFromTest(unittest.TestCase):
    def test_first_price(self):
        self.assertEqual(pr._hist_key(pr._apply_from([], 3000, ts(2026, 1, 1), 1)),
                         [(3000, ts(2026, 1, 1), None)])

    def test_empty_price_without_history_writes_nothing(self):
        self.assertEqual(pr._apply_from([], None, ts(2026, 1, 1), 1), [])

    def test_later_versions_are_replaced(self):
        old = [ver(2800, ts(2026, 9, 1), ts(2026, 10, 10)), ver(3000, ts(2026, 10, 10))]
        new = pr._apply_from(old, 3200, ts(2026, 10, 1), 1)
        self.assertEqual(pr._hist_key(new), [(2800, ts(2026, 9, 1), ts(2026, 10, 1)),
                                             (3200, ts(2026, 10, 1), None)])

    def test_same_price_as_before_extends_previous_version(self):
        old = [ver(2800, ts(2026, 9, 1), ts(2026, 10, 10)), ver(3000, ts(2026, 10, 10))]
        new = pr._apply_from(old, 2800, ts(2026, 10, 5), 1)
        self.assertEqual(pr._hist_key(new), [(2800, ts(2026, 9, 1), None)])

    def test_unchanged_reupload_is_noop(self):
        old = [ver(3000, ts(2026, 1, 1))]
        self.assertEqual(pr._hist_key(pr._apply_from(old, 3000, ts(2026, 1, 1), 1)), pr._hist_key(old))

    def test_earlier_date_moves_start(self):
        old = [ver(3000, ts(2026, 10, 10))]
        self.assertEqual(pr._hist_key(pr._apply_from(old, 3000, ts(2026, 1, 1), 1)),
                         [(3000, ts(2026, 1, 1), None)])

    def test_clear_price_from_date(self):
        old = [ver(3000, ts(2026, 1, 1))]
        self.assertEqual(pr._hist_key(pr._apply_from(old, None, ts(2026, 10, 1), 1)),
                         [(3000, ts(2026, 1, 1), ts(2026, 10, 1)), (None, ts(2026, 10, 1), None)])


class DateCellTest(unittest.TestCase):
    TODAY = dt.date(2026, 10, 12)

    def test_formats(self):
        for v in (dt.datetime(2026, 1, 1), dt.date(2026, 1, 1), "01.01.2026", "2026-01-01", " 01.01.26 "):
            with self.subTest(v=v):
                self.assertEqual(pr._date_cell(v, self.TODAY), dt.date(2026, 1, 1))
        self.assertEqual(pr._date_cell(46023, self.TODAY), dt.date(2026, 1, 1))   # дата Excel без формата
        for v in (None, "", " ", "—"):
            self.assertIsNone(pr._date_cell(v, self.TODAY))

    def test_rejects_future_and_garbage(self):
        for v in ("13.10.2026", "31.02.2026", "завтра", dt.date(2019, 12, 31)):
            with self.subTest(v=v), self.assertRaises(ValueError):
                pr._date_cell(v, self.TODAY)


def xlsx(rows, header=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header or ["Артикул", pr.COL_OZON, pr.COL_OZON_FROM, pr.COL_WB, pr.COL_WB_FROM,
                         pr.COL_YA, pr.COL_YA_FROM])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return UploadFile(filename="prices.xlsx", file=buf)


async def _collect(resp):
    return b"".join([chunk async for chunk in resp.body_iterator])


def hist(art="SKU-1", mp="Wildberries"):
    return [(r["base_price"], r["valid_from"].astimezone(MSK).date(),
             r["valid_to"].astimezone(MSK).date() if r["valid_to"] else None)
            for r in db.query_all("""SELECT * FROM catalog_base_prices_hist
                                      WHERE seller_article=%s AND marketplace=%s
                                      ORDER BY valid_from""", (art, mp))]


class ImportTest(unittest.TestCase):
    def setUp(self):
        prepare_local_db()
        daily_pi.migrate()

    def test_dated_import_export_roundtrip_and_period_boundary(self):
        r = pr.prices_import(xlsx([["SKU-1", None, None, 3000, "01.01.2026", None, None]]), user=ADMIN)
        self.assertEqual((r["prices_set"], r["dated"], r["history_changed"]), (1, 1, 1))
        self.assertEqual(hist(), [(3000, dt.date(2026, 1, 1), None)])
        self.assertEqual(db.query_one("""SELECT base_price FROM catalog_base_prices
            WHERE seller_article='SKU-1' AND marketplace='Wildberries'""")["base_price"], 3000)

        # С начала октября новая цена; в прайс-листе за сентябрь — старая (граница — сутки по Москве)
        pr.prices_import(xlsx([["SKU-1", None, None, 3200, dt.datetime(2026, 10, 1), None, None]]), user=ADMIN)
        self.assertEqual(hist(), [(3000, dt.date(2026, 1, 1), dt.date(2026, 10, 1)),
                                  (3200, dt.date(2026, 10, 1), None)])
        sept = {i["seller_article"]: i for i in pr.pricelist("2026-09-01", "2026-09-30", ADMIN)["items"]}
        self.assertEqual((sept["SKU-1"]["price_wb"], sept["SKU-1"]["from_wb"]), (3000, "2026-01-01"))

        # Выгрузка → та же загрузка без правок: история не меняется
        resp = pr.prices_export(None, None, False, "", ADMIN)
        body = asyncio.run(_collect(resp))
        wb = openpyxl.load_workbook(io.BytesIO(body))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:7], ("Артикул", pr.COL_OZON, pr.COL_OZON_FROM, pr.COL_WB,
                                       pr.COL_WB_FROM, pr.COL_YA, pr.COL_YA_FROM))
        row = next(r for r in rows if r[0] == "SKU-1")
        self.assertEqual((row[3], row[4].date()), (3200, dt.date(2026, 10, 1)))
        buf = io.BytesIO(body)
        r = pr.prices_import(UploadFile(filename="back.xlsx", file=buf), user=ADMIN)
        self.assertEqual(r["history_changed"], 0)
        self.assertEqual(len(hist()), 2)

    def test_old_file_without_dates_still_works(self):
        r = pr.prices_import(xlsx([["SKU-1", 1500, 2500, 900]],
                                  header=["Артикул", pr.COL_OZON, pr.COL_WB, pr.COL_YA]), user=ADMIN)
        self.assertEqual((r["prices_set"], r["dated"]), (3, 0))
        self.assertEqual([h[0] for h in hist("SKU-1", "Ozon")], [1500])
        # Пустые ячейки без цены — не «очищено»; очистка существующей цены — считается
        r = pr.prices_import(xlsx([["SKU-1", None, 2500, None], ["SKU-2", None, None, None]],
                                  header=["Артикул", pr.COL_OZON, pr.COL_WB, pr.COL_YA]), user=ADMIN)
        # SKU-1: Ozon и Yandex; SKU-2: WB 1000 из prepare_local_db; пустые без цены не считаются
        self.assertEqual(r["prices_cleared"], 3)

    def test_future_date_rejects_whole_file(self):
        future = (dt.datetime.now(MSK).date() + dt.timedelta(days=1)).strftime("%d.%m.%Y")
        with self.assertRaises(HTTPException) as e:
            pr.prices_import(xlsx([["SKU-1", 1000, "01.01.2026", 2000, future, None, None]]), user=ADMIN)
        self.assertIn("в будущем", e.exception.detail)
        self.assertEqual(hist("SKU-1", "Ozon"), [])

    def test_dated_price_feeds_daily_pi(self):
        with db.transaction() as tx:
            tx.execute("""INSERT INTO mp_price_daily (seller_article, marketplace, date, upload_price,
                              buyer_price, spp_pct, spp_source_date, upload_source, spp_is_estimated)
                          VALUES ('SKU-1','Wildberries','2026-09-18',1000,800,.2,'2026-09-18',
                                  'upload_wb_spp',false)""")
        pr.prices_import(xlsx([["SKU-1", None, None, 500, "01.09.2026", None, None]]), user=ADMIN)
        pi = db.query_one("SELECT wb_base_pi FROM price_index_daily WHERE seller_article='SKU-1' AND date='2026-09-18'")
        self.assertEqual(float(pi["wb_base_pi"]), 2.0)


if __name__ == "__main__":
    unittest.main()
