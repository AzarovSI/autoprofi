"""Workbook and read-only endpoint regression tests (no production database)."""
import io
import unittest
from unittest.mock import patch

import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app import auth
from app.price_export import index_workbook
from app.routers import prices_router as prices


def fixture():
    return [
        dict(seller_article="AGR-35", sample_name="Компрессор", l1="Продукция ECOM",
             l2="Компрессоры", l3="Автомобильные", price_wb=1000, price_ozon=1200,
             price_ya=1300, cost=555.25, cost_date="2026-09-01",
             wb_upload_price=1100, wb_spp_pct=.2, wb_buyer_price=880,
             oz_upload_price=1320, oz_spp_pct=.25, oz_buyer_price=990,
             wb_last_date="2026-09-18", oz_last_date="2026-09-19"),
        dict(seller_article="00123", sample_name="Накидка", l1="Продукция ТД АВТОПРОФИ",
             l2="Аксессуары", l3="Накидки", price_wb=500, price_ozon=600,
             cost=None, wb_upload_price=1000, wb_spp_pct=.4, wb_buyer_price=600,
             oz_upload_price=1100, oz_spp_pct=.4, oz_buyer_price=660),
        dict(seller_article="OUT", sample_name="Нет в наличии", l1=None,
             price_wb=0, price_ozon=None, cost=0, wb_upload_price=1000,
             wb_spp_pct=None, wb_buyer_price=None, oz_upload_price=1200,
             oz_spp_pct=None, oz_buyer_price=None),
    ]


def workbook(buf, values=False):
    if hasattr(buf, "seek"):
        buf.seek(0)
    return openpyxl.load_workbook(buf, data_only=values)


def rows(ws):
    return {str(row[0].value): row[0].row for row in ws if row[0].value}


class PriceExportTests(unittest.TestCase):
    def test_layout_numeric_formats_outline_and_cost_default_off(self):
        wb = workbook(index_workbook(fixture()))
        ws = wb.active
        self.assertEqual(ws.max_column, 12)
        self.assertEqual(ws.freeze_panes, "B6")
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)
        r = rows(ws)["AGR-35"]
        self.assertEqual(ws.row_dimensions[r].outlineLevel, 3)
        self.assertEqual(ws.cell(r, 5).number_format, "0.0%")
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "001F3B57")
        self.assertNotIn("Себестоимость", str([c.value for row in ws for c in row]))

    def test_pi_formulas_and_cached_values_match_screen(self):
        buf = index_workbook(fixture())
        formulas, values = workbook(buf).active, workbook(buf, True).active
        r = rows(values)["AGR-35"]
        self.assertEqual(formulas.cell(r, 3).data_type, "f")
        self.assertIn(f"D{r}/B{r}", formulas.cell(r, 3).value)
        self.assertAlmostEqual(values.cell(r, 3).value, 1.1)
        self.assertAlmostEqual(values.cell(r, 9).value, 1.1)
        self.assertAlmostEqual(values.cell(r, 7).value, 1.125)
        ar = rows(values)["00123"]
        self.assertIn(f"F{ar}/B{ar}", formulas.cell(ar, 3).value)
        self.assertAlmostEqual(values.cell(ar, 3).value, 1.2)
        self.assertAlmostEqual(values.cell(ar, 9).value, 1.1)
        self.assertEqual(formulas.cell(ar, 1).data_type, "s")

    def test_cost_enabled_and_shifted_formulas(self):
        buf = index_workbook(fixture(), show_cost=True)
        ws = workbook(buf, True).active
        self.assertEqual(ws.max_column, 13)
        self.assertEqual(ws["B4"].value, "Себестоимость, ₽")
        r = rows(ws)["AGR-35"]
        self.assertEqual(ws.cell(r, 2).value, 555.25)
        self.assertIn("2026-09-01", ws.cell(r, 2).comment.text)
        self.assertAlmostEqual(ws.cell(r, 8).value, 1.125)
        self.assertEqual(ws.cell(rows(ws)["OUT"], 2).value, 0)
        self.assertEqual(ws.cell(rows(ws)["00123"], 2).value, "–")

    def test_missing_prices_are_dashes_and_zero_denominator_safe(self):
        ws = workbook(index_workbook(fixture()), True).active
        r = rows(ws)["OUT"]
        for col in (3, 5, 6, 7, 8, 9, 11, 12):
            self.assertEqual(ws.cell(r, col).value, "–")
        self.assertEqual(ws.cell(r, 4).value, 1000)
        self.assertEqual(ws.cell(r, 10).value, 1200)

    def test_catalog_formula_injection_and_leading_zeros(self):
        data = fixture()
        data[0]["seller_article"] = '=HYPERLINK("https://example.invalid")'
        data[0]["l1"] = "=1+1"
        ws = workbook(index_workbook(data)).active
        for key in (data[0]["seller_article"], "=1+1", "00123"):
            self.assertEqual(ws.cell(rows(ws)[key], 1).data_type, "s")
        self.assertTrue(all(c.data_type != "f" for row in ws for c in row if c.column == 1))

    def test_estimated_spp_numeric_value_and_marker(self):
        data = fixture()
        data[0]["wb_spp_is_estimated"] = True
        ws = workbook(index_workbook(data)).active
        c = ws.cell(rows(ws)["AGR-35"], 5)
        self.assertEqual(c.value, .2)
        self.assertEqual(c.number_format, '"≈"0.0%')
        self.assertIn("оценочная", c.comment.text)

    def test_empty_report_valid_and_search_is_text(self):
        ws = workbook(index_workbook([], search="=1+1")).active
        self.assertEqual(ws["A6"].value, "Ничего не найдено")
        self.assertIn("=1+1", ws["A3"].value)
        self.assertEqual(ws["A3"].data_type, "s")

    def test_endpoint_authorization(self):
        app = FastAPI()
        app.include_router(prices.router)
        client = TestClient(app)
        for path in ("index_export", "export"):
            self.assertEqual(client.get("/api/prices/" + path).status_code, 401)

    def test_endpoint_period_search_and_no_cache(self):
        app = FastAPI()
        app.include_router(prices.router)
        app.dependency_overrides[auth.get_current_user] = lambda: {"id": 1}
        client = TestClient(app)
        data = {"items": fixture(), "period": {"date_from": "2026-09-01", "date_to": "2026-09-19"}}
        with patch.object(prices, "pricelist", return_value=data) as read:
            response = client.get("/api/prices/index_export?date_from=2026-09-01&date_to=2026-09-19&search=компр&show_cost=true")
        self.assertEqual(response.status_code, 200)
        read.assert_called_once_with(date_from="2026-09-01", date_to="2026-09-19", user={"id": 1})
        ws = workbook(io.BytesIO(response.content), True).active
        self.assertIn("AGR-35", rows(ws))
        self.assertNotIn("00123", rows(ws))
        self.assertEqual(ws.max_column, 13)
        self.assertIn("private", response.headers["cache-control"])
        self.assertIn("2026-09-19.xlsx", response.headers["content-disposition"])

    def test_base_export_cost_optional_import_layout_preserved(self):
        app = FastAPI()
        app.include_router(prices.router)
        app.dependency_overrides[auth.get_current_user] = lambda: {"id": 1}
        client = TestClient(app)
        with patch.object(prices, "pricelist", return_value={"items": fixture(), "period": None}):
            for flag, cols in (("false", 4), ("true", 5)):
                with self.subTest(flag=flag):
                    response = client.get("/api/prices/export?show_cost=" + flag + "&search=agr")
                    ws = workbook(io.BytesIO(response.content), True).active
                    self.assertEqual(ws.max_column, cols)
                    self.assertEqual(ws.max_row, 2)
                    self.assertEqual([c.value for c in ws[1]][:4],
                                     ["Артикул", "Базовая OZON", "Базовая WB", "Базовая Yandex"])
                    if cols == 5:
                        self.assertEqual(ws["E2"].value, 555.25)

    def test_validation(self):
        app = FastAPI()
        app.include_router(prices.router)
        app.dependency_overrides[auth.get_current_user] = lambda: {"id": 1}
        client = TestClient(app)
        for path in ("index_export", "export"):
            self.assertEqual(client.get(f"/api/prices/{path}?show_cost=invalid").status_code, 422)
            self.assertEqual(client.get(f"/api/prices/{path}?date_from=bad").status_code, 400)
            self.assertEqual(client.get(f"/api/prices/{path}?search=" + "a" * 501).status_code, 422)


if __name__ == "__main__":
    unittest.main()
