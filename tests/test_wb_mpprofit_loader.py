"""Загрузчик MPPROFIT WB: индекс заголовков и сверка до записи (без БД)."""
import os
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl

from app import loader_core as lc

# Раскладка как в отчёте MPPROFIT WB: id-колонки только в строке 4, над метриками —
# секции. Секции «Продажи» и «Налог» совпадают с именами метрик — ловушка индекса.
IDS = ["Брэнд", "Наименование", "Артикул продавца", "Штрихкод"]
METRICS = [  # (секция строки 4 или None, метрика строки 5)
    ("Продажи", "Заказы шт."), (None, "Продажи, шт."), (None, "Возвраты шт."),
    (None, "Отмены шт."), (None, "Продажи"), (None, "Возвраты"), (None, "Выручка"),
    (None, "Итого продаж шт."), ("Себестоимость", "Себестоимость"), ("Удержания WB", "Комиссия"),
    (None, "Прочие удержания"), (None, "Включая продвижение"), (None, "Все удержания"),
    ("Налог", "Налоговая база"), (None, "Налог"), ("Финансы", "Выплата на р/с"), (None, "Прибыль"),
    ("Фин.результат", "Вся реклама"), (None, "Баланс"), (None, "Промо бонусы"),
    (None, "Выплата на р/с (с рекл.)"), (None, "Прибыль (с рекл.)"),
]


def item(art, revenue, holds, balance, bonus=0.0, orders=10):
    payout = revenue - holds
    return {"Артикул продавца": art, "Наименование": "Товар " + art, "Штрихкод": "46" + art,
            "Заказы шт.": orders, "Продажи, шт.": 5, "Возвраты шт.": 0, "Отмены шт.": 0,
            "Продажи": revenue, "Возвраты": 0, "Выручка": revenue, "Итого продаж шт.": 5,
            "Себестоимость": 100, "Комиссия": holds, "Все удержания": holds,
            "Налоговая база": revenue, "Налог": 0, "Выплата на р/с": payout,
            "Прибыль": payout - 100, "Вся реклама": balance + bonus, "Баланс": balance,
            "Промо бонусы": bonus, "Выплата на р/с (с рекл.)": payout - balance,
            "Прибыль (с рекл.)": payout - 100 - balance}


def holds_row(other, promo):
    # Строка «Удержания WB»: в «(с рекл.)» удержанная WB реклама возвращена обратно.
    return {"Наименование": "Удержания WB", "Прочие удержания": other,
            "Включая продвижение": promo, "Все удержания": other,
            "Выплата на р/с": -other, "Прибыль": -other,
            "Выплата на р/с (с рекл.)": -other + promo, "Прибыль (с рекл.)": -other + promo}


def make_ws(rows, drop=(), total_override=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(2, 3, "Период: 2026-09-01 - 2026-09-30")
    cols = IDS + [m for _, m in METRICS if m not in drop]
    sections = dict((m, s) for s, m in METRICS)
    for c, name in enumerate(cols, 1):
        if name in IDS:
            ws.cell(4, c, name)
        else:
            if sections.get(name):
                ws.cell(4, c, sections[name])
            ws.cell(5, c, name)
    ws.cell(6, 1, "Итого")
    for r, row in enumerate(rows, 7):
        for c, name in enumerate(cols, 1):
            if name in row:
                ws.cell(r, c, row[name])
    for c, name in enumerate(cols, 1):
        if name not in IDS:
            ws.cell(6, c, sum(row.get(name, 0) or 0 for row in rows))
    for name, v in (total_override or {}).items():
        ws.cell(6, cols.index(name) + 1, v)
    return wb, ws


ROWS = [item("A-1", 1000, 300, 50, 2), item("B-2", 2000, 500, 120), holds_row(700, 650)]


class WbHeaderIndexTest(unittest.TestCase):
    def test_metric_row_wins_over_section(self):
        _, ws = make_ws(ROWS)
        idx = lc.build_colidx(ws, metric_row_first=True)
        self.assertEqual(ws.cell(5, idx["Продажи"]).value, "Продажи")
        self.assertEqual(ws.cell(5, idx["Налог"]).value, "Налог")
        self.assertEqual(ws.cell(4, idx["Артикул продавца"]).value, "Артикул продавца")

    def test_old_order_kept_by_default(self):
        _, ws = make_ws(ROWS)
        idx = lc.build_colidx(ws)
        self.assertEqual(ws.cell(5, idx["Продажи"]).value, "Заказы шт.")


class WbCheckTest(unittest.TestCase):
    def check(self, rows=ROWS, **kw):
        _, ws = make_ws(rows, **kw)
        return lc.check_wb_report(ws, lc.build_colidx(ws, metric_row_first=True))

    def test_consistent_report_passes(self):
        self.assertEqual(self.check(), ([], []))

    def test_total_mismatch_rejected(self):
        errors, _ = self.check(total_override={"Выручка": 2990})
        self.assertEqual(len(errors), 1)
        self.assertIn("«Выручка»", errors[0])
        self.assertIn("2 990,00", errors[0])

    def test_payout_identity_rejected(self):
        rows = [dict(ROWS[0], **{"Выплата на р/с": 1}), ROWS[1], ROWS[2]]
        errors, _ = self.check(rows)
        self.assertTrue(any("«Выручка» − «Все удержания»" in e and "A-1" in e for e in errors))

    def test_payout_with_ads_identity_rejected_on_item(self):
        rows = [ROWS[0], dict(ROWS[1], **{"Выплата на р/с (с рекл.)": 1500}), ROWS[2]]
        errors, _ = self.check(rows)
        self.assertTrue(any("«Баланс»" in e and "B-2" in e for e in errors))

    def test_old_report_without_new_columns_warns(self):
        errors, warnings = self.check(drop=lc.WB_OPTIONAL)
        self.assertEqual(errors, [])
        self.assertEqual(len(warnings), 1)
        self.assertIn("«Баланс»", warnings[0])


class WbLoadRejectTest(unittest.TestCase):
    def test_failed_check_writes_only_journal(self):
        wb, _ = make_ws(ROWS, total_override={"Выплата на р/с": 1})
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(path)
            conn = MagicMock()
            cur = conn.cursor.return_value.__enter__.return_value
            res = lc.load_report(conn, path, "wb")
        finally:
            os.remove(path)
        self.assertFalse(res["ok"])
        self.assertIn("ОШИБКА СВЕРКИ", res["error"])
        sqls = [c.args[0] for c in cur.execute.call_args_list]
        self.assertEqual(len(sqls), 1)
        self.assertIn("INSERT INTO report_uploads", sqls[0])
        self.assertEqual(cur.execute.call_args.args[1][9], "ОШИБКА СВЕРКИ")


if __name__ == "__main__":
    unittest.main()
