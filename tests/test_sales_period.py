# -*- coding: utf-8 -*-
"""«Продажи, шт.»: фильтр периода по месяцам (с 15.10.2026, решение владельца).

Период может переходить через границу года; шаблон плана — строго по месяцам периода;
загрузка берёт месяцы и год из шапки файла и перезаписывает только их.
Тесты с записью — только на локальной wb_price_test.
"""
import asyncio
import io
import unittest
from unittest.mock import patch

import openpyxl
from fastapi import HTTPException, UploadFile

from test_wb_prices import prepare_local_db
from app import db
from app.routers import sales_router as sr

ADMIN = {"id": 1, "role": "admin", "is_admin": True}


def plan_file(header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return UploadFile(filename="plan.xlsx", file=buf)


async def _collect(resp):
    return b"".join([chunk async for chunk in resp.body_iterator])


class PeriodTest(unittest.TestCase):
    def test_cross_year_period(self):
        self.assertEqual(sr._period_months("2026-11", "2027-02"),
                         [(2026, 11), (2026, 12), (2027, 1), (2027, 2)])

    def test_reversed_bounds_are_swapped(self):
        self.assertEqual(sr._period_months("2027-02", "2026-12"),
                         [(2026, 12), (2027, 1), (2027, 2)])

    def test_old_year_param_means_whole_year(self):
        self.assertEqual(sr._period_months(year=2025), [(2025, m) for m in range(1, 13)])

    def test_too_long_and_bad_values_rejected(self):
        with self.assertRaises(HTTPException):
            sr._period_months("2024-01", "2027-01")
        for bad in ("2026-13", "2026", "abc"):
            with self.assertRaises(HTTPException):
                sr._period_months(bad, "2026-12")

    def test_labels_and_tags(self):
        year = sr._period_months("2026-01", "2026-12")
        cross = sr._period_months("2026-11", "2027-02")
        one = sr._period_months("2026-10", "2026-10")
        self.assertEqual(sr._period_label(year), "2026 г.")
        self.assertEqual(sr._period_label(cross), "ноябрь 2026 – февраль 2027")
        self.assertEqual(sr._period_label(one), "октябрь 2026")
        self.assertEqual(sr._period_tag(year), "2026")
        self.assertEqual(sr._period_tag(cross), "2026-11_2027-02")
        self.assertEqual(sr._period_tag(one), "2026-10")


class HeaderTest(unittest.TestCase):
    def test_months_with_year_across_years(self):
        cols = sr._plan_header_months(
            ["Артикул", "Уровень 1", "Текущий статус", "Декабрь 2026", "Январь 2027"])
        self.assertEqual(cols, {3: (2026, 12), 4: (2027, 1)})

    def test_month_without_year_rejected(self):
        with self.assertRaises(HTTPException) as cm:
            sr._plan_header_months(["Артикул", "Январь", "Февраль"])
        self.assertIn("нет года", cm.exception.detail)

    def test_duplicate_month_rejected(self):
        with self.assertRaises(HTTPException):
            sr._plan_header_months(["Артикул", "Январь 2026", "Январь 2026"])

    def test_no_month_columns_rejected(self):
        with self.assertRaises(HTTPException):
            sr._plan_header_months(["Артикул", "Количество"])


class ImportTest(unittest.TestCase):
    """Загрузка файла за часть года не трогает план остальных месяцев."""

    def setUp(self):
        prepare_local_db()
        db.execute("TRUNCATE sales_plan")
        rows = [(a, "Ozon", 2026, m, 100, False, 100)
                for a in ("SKU-1", "SKU-2") for m in range(1, 13)]
        rows += [("SKU-1", "Ozon", 2027, 1, 50, True, None),
                 ("SKU-1", "Wildberries", 2026, 10, 7, False, 7)]
        db.execute_values(
            "INSERT INTO sales_plan (seller_article, marketplace, year, month, plan_qty, is_manual, prev_qty) VALUES %s",
            rows)
        db.execute("TRUNCATE catalog_marketplace")
        db.execute_values(
            "INSERT INTO catalog_marketplace (seller_article, marketplace) VALUES %s",
            [("SKU-1", "Ozon"), ("SKU-2", "Ozon"), ("SKU-1", "Wildberries"), ("SKU-3", "Wildberries")])

    def plan(self, art, mp="Ozon"):
        return {(r["year"], r["month"]): float(r["plan_qty"]) for r in db.query_all(
            "SELECT year, month, plan_qty FROM sales_plan WHERE seller_article=%s AND marketplace=%s",
            (art, mp))}

    def test_quarter_file_keeps_other_months(self):
        f = plan_file(["Артикул", "Текущий статус", "Октябрь 2026", "Ноябрь 2026", "Декабрь 2026"],
                      [["SKU-1", "CORE", 300, None, 320]])
        r = sr.plan_import(file=f, marketplace="ozon", user=ADMIN)
        self.assertEqual((r["month_from"], r["month_to"], r["months"]), ("2026-10", "2026-12", 3))
        p = self.plan("SKU-1")
        for m in range(1, 10):
            self.assertEqual(p[(2026, m)], 100)       # январь–сентябрь не тронуты
        self.assertEqual(p[(2026, 10)], 300)
        self.assertNotIn((2026, 11), p)                # пустая ячейка очищает свой месяц
        self.assertEqual(p[(2026, 12)], 320)
        self.assertEqual(p[(2027, 1)], 50)             # другой год не тронут
        self.assertEqual(self.plan("SKU-2")[(2026, 10)], 100)  # артикула нет в файле
        self.assertEqual(self.plan("SKU-1", "Wildberries")[(2026, 10)], 7)  # другой МП

    def test_cross_year_file(self):
        f = plan_file(["Артикул", "Декабрь 2026", "Январь 2027"], [["SKU-1", 11, 22]])
        sr.plan_import(file=f, marketplace="ozon", user=ADMIN)
        p = self.plan("SKU-1")
        self.assertEqual((p[(2026, 11)], p[(2026, 12)], p[(2027, 1)]), (100, 11, 22))
        row = db.query_one("""SELECT is_manual, prev_qty FROM sales_plan
                              WHERE seller_article='SKU-1' AND marketplace='Ozon' AND year=2027 AND month=1""")
        self.assertEqual((row["is_manual"], float(row["prev_qty"])), (False, 22))

    def test_bad_header_writes_nothing(self):
        f = plan_file(["Артикул", "Октябрь", "Ноябрь"], [["SKU-1", 1, 2]])
        with self.assertRaises(HTTPException):
            sr.plan_import(file=f, marketplace="ozon", user=ADMIN)
        self.assertEqual(self.plan("SKU-1")[(2026, 10)], 100)

    def test_article_not_in_catalog_rejects_whole_file(self):
        # SKU-3 есть в справочнике только у WB — для плана Ozon он неизвестен
        f = plan_file(["Артикул", "Октябрь 2026"], [["SKU-1", 300], ["SKU-3", 5], ["NEW-1", 7]])
        with self.assertRaises(HTTPException) as cm:
            sr.plan_import(file=f, marketplace="ozon", user=ADMIN)
        self.assertEqual(cm.exception.status_code, 400)
        self.assertIn("NEW-1, SKU-3", cm.exception.detail)
        self.assertNotIn("SKU-1", cm.exception.detail)
        self.assertEqual(self.plan("SKU-1")[(2026, 10)], 100)   # известный артикул тоже не записан
        self.assertEqual(self.plan("SKU-3"), {})
        self.assertEqual(self.plan("NEW-1"), {})

    def test_article_matched_after_canon(self):
        f = plan_file(["Артикул", "Октябрь 2026"], [["  sku-1 ", 300]])
        sr.plan_import(file=f, marketplace="ozon", user=ADMIN)
        self.assertEqual(self.plan("SKU-1")[(2026, 10)], 300)

    def test_plan_cell_needs_catalog(self):
        cell = sr.PlanCellIn(seller_article="SKU-3", marketplace="ozon", year=2026, month=10, plan_qty=5)
        with self.assertRaises(HTTPException) as cm:
            sr.plan_cell(cell, user=ADMIN)
        self.assertIn("SKU-3", cm.exception.detail)
        self.assertEqual(self.plan("SKU-3"), {})
        cell = sr.PlanCellIn(seller_article="SKU-3", marketplace="wb", year=2026, month=10, plan_qty=5)
        sr.plan_cell(cell, user=ADMIN)
        self.assertEqual(self.plan("SKU-3", "Wildberries")[(2026, 10)], 5)


class ExportTest(unittest.TestCase):
    def test_template_has_only_period_months_and_reimports(self):
        def query_all(sql, params=None):
            if "FROM sales_plan" in sql and "plan_qty" in sql:
                return [{"seller_article": "SKU-1", "year": 2026, "month": 12, "plan_qty": 5, "is_manual": True},
                        {"seller_article": "SKU-1", "year": 2027, "month": 1, "plan_qty": 6, "is_manual": False}]
            if "FROM catalog_items" in sql:
                return []
            return [{"seller_article": "SKU-1", "status": "CORE"}]

        with patch.object(sr.db, "query_all", side_effect=query_all):
            resp = sr.plan_export(marketplace="ozon", month_from="2026-12", month_to="2027-01",
                                  year=None, user=ADMIN)
        self.assertIn("plan_ozon_2026-12_2027-01.xlsx", resp.headers["content-disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(asyncio.run(_collect(resp)))).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][5:], ("Декабрь 2026", "Январь 2027"))
        self.assertEqual(rows[1][0], "SKU-1")
        self.assertEqual(rows[1][5:], (5, 6))
        self.assertEqual(ws.cell(row=2, column=6).font.color.rgb, "0016A360")  # ручная — зелёная
        self.assertEqual(sr._plan_header_months(rows[0]), {5: (2026, 12), 6: (2027, 1)})


class ComputeTest(unittest.TestCase):
    """Расчёт дерева: ячейки идут по месяцам периода, YoY — тот же месяц годом раньше."""

    def test_cross_year_cells_and_yoy(self):
        def query_all(sql, params=None):
            if "FROM fact_monthly" in sql and "SUM(sales_qty * " in sql:
                return [{"seller_article": "A", "marketplace": "Ozon", "year": 2026, "month": 12, "qty": 10, "rub": 1000},
                        {"seller_article": "A", "marketplace": "Ozon", "year": 2027, "month": 1, "qty": 4, "rub": 400}]
            if "FROM fact_monthly" in sql and "SUM(sales_qty) AS qty" in sql:
                self.assertEqual(params[:2], (202512, 202601))  # окно годом раньше
                return [{"seller_article": "A", "marketplace": "Ozon", "year": 2025, "month": 12, "qty": 8},
                        {"seller_article": "A", "marketplace": "Ozon", "year": 2026, "month": 1, "qty": 3}]
            if "FROM sales_plan" in sql:
                return [{"seller_article": "A", "marketplace": "Ozon", "year": 2027, "month": 1,
                         "plan_qty": 20, "is_manual": True, "prev_qty": 15}]
            if "FROM catalog_items" in sql:
                return [{"seller_article": "A", "category_l1": "L1", "category_l2": "L2", "category_l3": "L3"}]
            if "SELECT DISTINCT year" in sql:
                return [{"year": 2026}, {"year": 2027}]
            return []

        with patch.object(sr.db, "query_all", side_effect=query_all):
            d = sr._compute_plan_fact("ozon", sr._period_months("2026-12", "2027-01"))
        self.assertEqual([m["key"] for m in d["months"]], ["2026-12", "2027-01"])
        self.assertEqual((d["month_from"], d["month_to"]), ("2026-12", "2027-01"))
        leaf = d["tree"]["children"][0]["children"][0]["children"][0]["children"][0]
        dec, jan = leaf["months"]
        self.assertEqual((dec["fact"], dec["fact_prev"], dec["plan"]), (10, 8, 0))
        self.assertEqual((jan["fact"], jan["fact_prev"], jan["plan"]), (4, 3, 20))
        self.assertEqual((jan["is_manual"], jan["prev_qty"]), (True, 15))
        self.assertEqual([m["plan"] for m in d["tree"]["months"]], [0, 20])


if __name__ == "__main__":
    unittest.main()
