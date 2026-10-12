"""Права доступа: каталог функций = xlsx владельца, миграция ролей. Только локальная wb_price_test."""
from pathlib import Path
import unittest

import openpyxl
import psycopg

from test_wb_prices import prepare_local_db
from app import auth, db, permissions

XLSX = Path(__file__).resolve().parent.parent / "docs" / "access" / "roles_matrix_20261012.xlsx"


def matrix():
    """(роли, строки функций, площадки ролей) с листов xlsx."""
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Права доступа"]
    roles = [c.value for c in ws[4]][7:15]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if isinstance(r[0], int)]
    limits = {r[2]: r[4:12] for r in wb["Ограничения данных"].iter_rows(min_row=5, values_only=True)
              if r[0]}
    names = {"Wildberries": "Видит данные Wildberries", "Ozon": "Видит данные Ozon",
             "Yandex": "Видит данные Яндекс.Маркета"}
    mps = [[mp for mp, label in names.items() if limits[label][i] == "✓"] for i in range(len(roles))]
    return roles, rows, mps


def reset_roles():
    """Убрать таблицы ролей, чтобы миграция выполнилась заново."""
    with psycopg.connect(db.DB_DSN, autocommit=True) as conn:
        conn.execute("ALTER TABLE app_users DROP COLUMN IF EXISTS role_id")
        conn.execute("DROP TABLE IF EXISTS app_role_permissions, app_roles")
        if conn.execute("SELECT to_regclass('app_schema_migrations')").fetchone()[0]:
            conn.execute("DELETE FROM app_schema_migrations WHERE name=%s", (permissions.MIGRATION,))


class CatalogTest(unittest.TestCase):
    def test_catalog_matches_xlsx(self):
        _, rows, _ = matrix()
        self.assertEqual(len(permissions.FUNCTIONS), 99)
        self.assertEqual(len(permissions.CODES), 99)
        for f, r in zip(permissions.FUNCTIONS, rows):
            with self.subTest(n=f[0]):
                self.assertEqual((f[0], *f[2:]), tuple(r[:6]))


class MigrationTest(unittest.TestCase):
    def setUp(self):
        prepare_local_db()
        reset_roles()
        permissions.migrate()

    def tearDown(self):
        reset_roles()

    def test_roles_and_permissions_match_xlsx(self):
        roles, rows, mps = matrix()
        saved = db.query_all("SELECT id, name, is_admin, marketplaces FROM app_roles ORDER BY id")
        self.assertEqual([r["name"] for r in saved], roles)
        self.assertEqual([r["is_admin"] for r in saved], [True] + [False] * 7)
        self.assertEqual([r["marketplaces"] for r in saved], mps)
        for i, role in enumerate(saved):
            with self.subTest(role=role["name"]):
                got = {r["permission"] for r in db.query_all(
                    "SELECT permission FROM app_role_permissions WHERE role_id=%s", (role["id"],))}
                expected = set() if role["is_admin"] else {
                    f[1] for f, r in zip(permissions.FUNCTIONS, rows) if r[7 + i] == "✓"}
                self.assertEqual(got, expected)
                self.assertLessEqual(got, permissions.CODES)

    def test_owner_gets_admin_role_others_untouched(self):
        reset_roles()
        owner = next(iter(auth.OWNER_EMAILS))
        db.execute("INSERT INTO app_users (id, email, display_name, role) VALUES (50, %s, 'Владелец', 'admin')",
                   (owner.upper(),))
        permissions.migrate()
        rows = {r["id"]: r for r in db.query_all(
            "SELECT u.id, r.is_admin FROM app_users u LEFT JOIN app_roles r ON r.id = u.role_id")}
        self.assertTrue(rows[50]["is_admin"])
        self.assertIsNone(rows[1]["is_admin"])  # пользователь из prepare_local_db — без роли
        db.execute("DELETE FROM app_users WHERE id = 50")

    def test_second_run_changes_nothing(self):
        db.execute("UPDATE app_roles SET name='РОП (переименована)' WHERE name='РОП'")
        permissions.migrate()
        self.assertEqual(db.query_one("SELECT count(*) AS n FROM app_roles")["n"], 8)
        self.assertTrue(db.query_one("SELECT 1 AS x FROM app_roles WHERE name='РОП (переименована)'"))

    def test_new_role_gets_next_id(self):
        row = db.query_one("INSERT INTO app_roles (name) VALUES ('Новая') RETURNING id, marketplaces")
        self.assertEqual((row["id"], row["marketplaces"]), (9, []))

    def test_only_one_admin_role(self):
        with self.assertRaises(psycopg.errors.UniqueViolation):
            db.execute("INSERT INTO app_roles (name, is_admin) VALUES ('Второй админ', true)")

    def test_unknown_marketplace_rejected(self):
        with self.assertRaises(psycopg.errors.CheckViolation):
            db.execute("INSERT INTO app_roles (name, marketplaces) VALUES ('X', '{Avito}')")

    def test_role_with_users_cannot_be_deleted(self):
        db.execute("UPDATE app_users SET role_id=(SELECT id FROM app_roles WHERE name='Ассистент')")
        with self.assertRaises(psycopg.errors.RestrictViolation):
            db.execute("DELETE FROM app_roles WHERE name='Ассистент'")

    def test_deleting_role_removes_its_permissions(self):
        role_id = db.query_one("SELECT id FROM app_roles WHERE name='Маркетинг'")["id"]
        db.execute("DELETE FROM app_roles WHERE id=%s", (role_id,))
        self.assertIsNone(db.query_one(
            "SELECT 1 AS x FROM app_role_permissions WHERE role_id=%s", (role_id,)))


if __name__ == "__main__":
    unittest.main()
