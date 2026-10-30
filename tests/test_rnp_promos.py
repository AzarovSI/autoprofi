"""Акции товаров в РНП (rnp_promos). Только локальная wb_price_test."""
import datetime
import io
import unittest

import openpyxl

import psycopg
from fastapi.testclient import TestClient

from test_wb_prices import prepare_local_db
from test_access_roles import reset_roles
from app import auth, db, permissions
from app.main import app
from app.routers import rnp_promos_router

D = datetime.date
TEMPLATE = [b""]  # шаблон, выгруженный в тесте: заполняется и загружается обратно


def reset_promos():
    with psycopg.connect(db.DB_DSN, autocommit=True) as conn:
        conn.execute("DROP TABLE IF EXISTS rnp_promo_events, rnp_promos")
        if conn.execute("SELECT to_regclass('app_schema_migrations')").fetchone()[0]:
            conn.execute("DELETE FROM app_schema_migrations WHERE name = ANY(%s)", (rnp_promos_router.MIGRATIONS,))


class RnpPromosTest(unittest.TestCase):
    def setUp(self):
        prepare_local_db()
        reset_roles()
        permissions.migrate()
        reset_promos()
        # «Менеджер WB+YA» владелец делает копией «Менеджер WB» с площадками WB и Яндекс.
        db.execute("INSERT INTO app_roles (name, is_admin, marketplaces) "
                   "VALUES ('Менеджер WB+YA', false, '{Wildberries,Yandex}')")
        db.execute("INSERT INTO app_role_permissions (role_id, permission) "
                   "SELECT (SELECT id FROM app_roles WHERE name = 'Менеджер WB+YA'), permission "
                   "FROM app_role_permissions WHERE role_id = (SELECT id FROM app_roles WHERE name = 'Менеджер WB')")
        rnp_promos_router.migrate()
        rnp_promos_router.migrate()  # повторный запуск ничего не меняет
        ids = {r["name"]: r["id"] for r in db.query_all("SELECT id, name FROM app_roles")}
        db.execute("INSERT INTO app_users (id, email, display_name, role, role_id) VALUES "
                   "(101, 'admin@local.test', 'Админ', 'admin', %s), (102, 'mkt@local.test', 'Маркетолог', 'user', %s), "
                   "(103, 'ast@local.test', 'Ассистент', 'user', %s), (104, 'wb@local.test', 'Менеджер WB', 'user', %s), "
                   "(105, 'wb2@local.test', 'Второй менеджер', 'user', %s), "
                   "(106, 'oz@local.test', 'Менеджер OZ', 'user', %s), (107, 'pr@local.test', 'Продакт', 'user', %s)",
                   (ids["Администратор"], ids["Маркетинг"], ids["Ассистент"], ids["Менеджер WB"], ids["Менеджер WB"],
                    ids["Менеджер OZ"], ids["Продакт"]))
        db.execute("INSERT INTO catalog_marketplace (seller_article, marketplace) VALUES "
                   "('SKU-1', 'Wildberries'), ('SKU-2', 'Wildberries'), ('SKU-3', 'Wildberries'), "
                   "('EDP24-20K', 'Wildberries'), ('00123', 'Wildberries'), ('SKU-1', 'Ozon')")
        self.client = TestClient(app, raise_server_exceptions=True)

    def tearDown(self):
        reset_promos()
        db.execute("DELETE FROM app_users WHERE id >= 100")
        db.execute("TRUNCATE catalog_marketplace")
        reset_roles()

    def call(self, uid, method, path, **kw):
        token = auth.create_token({"id": uid, "email": ""})
        return self.client.request(method, path, headers={"Authorization": f"Bearer {token}"}, **kw)

    def add(self, uid=104, mp="wb", sa="SKU-1", name="Осенний драйв", d_from="2026-10-10", plan="2026-10-20"):
        return self.call(uid, "POST", "/api/rnp-promos", json={
            "marketplace": mp, "seller_article": sa, "name": name, "date_from": d_from, "date_to_plan": plan})

    def listed(self, mp, uid=104):
        r = self.call(uid, "GET", f"/api/rnp-promos?marketplace={mp}&date_from=2026-10-01&date_to=2026-10-31")
        self.assertEqual(r.status_code, 200, r.text)
        return [(p["seller_article"], p["name"], p["date_from"], p["date_to"]) for p in r.json()["promos"]]

    def test_rights_by_role(self):
        """Решения владельца 23.10 и 30.10.2026: акции — ролям с правами на комментарии, кроме Продакта
        (только смотрит); площадки акций — по площадкам роли; загрузка и журнал — им же."""
        rows = db.query_all("SELECT r.name, p.permission FROM app_role_permissions p "
                            "JOIN app_roles r ON r.id = p.role_id WHERE p.permission LIKE 'rnp.promos.%%'")
        by_role = {}
        for r in rows:
            by_role.setdefault(r["name"], set()).add(r["permission"].removeprefix("rnp.promos."))
        acts = {"view", "add", "edit_own", "import", "journal"}
        self.assertEqual(by_role, {
            "РОП": acts | {"edit_any", "mp.wb", "mp.ozon", "mp.ya"},
            "Маркетинг": acts | {"mp.wb", "mp.ozon", "mp.ya"},
            "Менеджер WB": acts | {"mp.wb"},
            "Менеджер OZ": acts | {"mp.ozon"},
            "Менеджер YA": acts | {"mp.ya"},
            "Менеджер WB+YA": acts | {"mp.wb", "mp.ya"},
            "Продакт": {"view"},
            "Ассистент": {"view"},
        })

    def test_marketplace_right(self):
        """Менеджер WB видит акции Ozon, но открыть, изменить или отменить их не может."""
        self.assertEqual(self.add(uid=104, mp="ozon").status_code, 403)
        self.assertEqual(self.add(uid=107).status_code, 403)  # Продакт — только смотрит
        pid = self.add(uid=106, mp="ozon").json()["id"]
        self.assertEqual([p[0] for p in self.listed("ozon", uid=104)], ["SKU-1"])
        self.assertEqual(self.call(104, "POST", f"/api/rnp-promos/{pid}/cancel").status_code, 403)
        self.assertEqual(self.call(102, "POST", f"/api/rnp-promos/{pid}/close",
                                   json={"date_to_fact": "2026-10-12"}).status_code, 403)  # чужая без «чужую»
        self.assertEqual(self.call(106, "POST", f"/api/rnp-promos/{pid}/cancel").status_code, 200)

    @staticmethod
    def workbook(mp="Wildberries", name="Распродажа 11.11", d_from=D(2026, 11, 5), d_to=D(2026, 11, 15),
                 articles=()):
        wb = openpyxl.load_workbook(io.BytesIO(TEMPLATE[0]))
        ws = wb.active
        ws["B4"], ws["B5"], ws["B6"], ws["B7"] = mp, name, d_from, d_to
        for i, a in enumerate(articles):
            ws.cell(10 + i, 1, a)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def upload(self, content, uid=104, fname="Акция 11.11.xlsx"):
        return self.call(uid, "POST", "/api/rnp-promos/import", files={"file": (fname, content)})

    def test_template_is_empty_header(self):
        r = self.call(104, "GET", "/api/rnp-promos/template")
        self.assertEqual(r.status_code, 200, r.text)
        TEMPLATE[0] = r.content
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([ws[f"A{i}"].value for i in range(4, 8)],
                         ["Площадка", "Название акции", "Вход в акцию", "Плановый выход"])
        self.assertEqual(ws["A9"].value, "Артикул продавца")
        self.assertIsNone(ws["A10"].value)
        self.assertEqual(self.call(103, "GET", "/api/rnp-promos/template").status_code, 403)
        self.assertEqual(self.call(107, "GET", "/api/rnp-promos/template").status_code, 403)

    def test_import(self):
        TEMPLATE[0] = self.call(104, "GET", "/api/rnp-promos/template").content
        self.add(sa="SKU-3", name="Осенний драйв", d_from="2026-10-25", plan="2026-11-07")
        content = self.workbook(articles=["SKU-1", " edp24-20К ", "TT-902J", "SKU-3", "sku-1", None, 123, "00123"])
        r = self.upload(content)
        self.assertEqual(r.status_code, 200, r.text)
        j = r.json()
        self.assertEqual((j["marketplace"], j["name"], j["days"], j["rows"]), ("Wildberries", "Распродажа 11.11", 11, 7))
        self.assertEqual(j["added"], ["SKU-1", "edp24-20К", "00123"])
        self.assertEqual([(e["row"], e["article"], e["reason"]) for e in j["errors"]],
                         [(12, "TT-902J", "not_found"), (13, "SKU-3", "overlap"), (14, "sku-1", "duplicate"),
                          (16, "123", "not_found")])
        self.assertIn("Осенний драйв", j["errors"][1]["detail"])
        jr = self.call(102, "GET", "/api/rnp-promos/journal?marketplace=wb").json()["promos"]
        loaded = {p["seller_article"]: p for p in jr if p["name"] == "Распродажа 11.11"}
        self.assertEqual(sorted(loaded), ["00123", "EDP24-20K", "SKU-1"])  # артикул — к единому виду
        self.assertEqual({p["source_file"] for p in loaded.values()}, {"Акция 11.11.xlsx"})
        ev = self.call(102, "GET", f"/api/rnp-promos/{loaded['SKU-1']['id']}/events").json()["events"]
        self.assertEqual((ev[0]["action"], ev[0]["changes"]["file"], ev[0]["user_name"]),
                         ("create", "Акция 11.11.xlsx", "Менеджер WB"))
        # Повторная загрузка того же файла ничего не задваивает.
        j = self.upload(content).json()
        self.assertEqual((j["added"], [e["article"] for e in j["exists"]]), ([], ["SKU-1", "edp24-20К", "00123"]))
        self.assertEqual(len(self.call(102, "GET", "/api/rnp-promos/journal?marketplace=wb").json()["promos"]), 4)

    def test_import_rejects_bad_header_and_foreign_marketplace(self):
        TEMPLATE[0] = self.call(104, "GET", "/api/rnp-promos/template").content
        cases = [
            (dict(mp=None), "не выбрана площадка"),
            (dict(mp="Авито"), "неизвестная площадка «Авито»"),
            (dict(name=" "), "не указано название акции"),
            (dict(d_from="31.11.2026"), "вход в акцию «31.11.2026»"),
            (dict(d_to=D(2026, 11, 1)), "плановый выход раньше входа"),
            (dict(), "нет ни одного артикула"),
        ]
        for kw, msg in cases:
            with self.subTest(msg=msg):
                r = self.upload(self.workbook(articles=["SKU-1"] if kw else [], **kw))
                self.assertEqual(r.status_code, 422, r.text)
                self.assertIn(msg, r.json()["detail"])
        self.assertEqual(self.upload(b"not excel").status_code, 422)
        # Площадка не своя — файл не загружается целиком; строковые даты и «OZON» из списка читаются.
        ozon = self.workbook(mp="OZON", d_from="05.11.2026", d_to="2026-11-15", articles=["SKU-1"])
        self.assertEqual(self.upload(ozon).status_code, 403)
        r = self.upload(ozon, uid=106)
        self.assertEqual((r.status_code, r.json()["added"], r.json()["days"]), (200, ["SKU-1"], 11))
        self.assertEqual(self.upload(ozon, uid=103).status_code, 403)  # Ассистент: нет загрузки
        self.assertEqual(db.query_one("SELECT count(*) AS n FROM rnp_promos")["n"], 1)

    def test_create_list_and_one_promo_at_a_time(self):
        r = self.add()
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual((r.json()["marketplace"], r.json()["days"]), ("Wildberries", 11))
        self.assertEqual(self.listed("wb"), [("SKU-1", "Осенний драйв", "2026-10-10", "2026-10-20")])
        self.assertEqual(self.listed("ozon"), [])  # своя акция у каждой площадки
        # Пересечение — нельзя; на другой площадке и сразу после окончания — можно.
        r = self.add(name="Вторая", d_from="2026-10-20", plan="2026-10-25")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Осенний драйв", r.json()["detail"])
        self.assertEqual(self.add(uid=106, mp="ozon", name="Ozon", d_from="2026-10-15", plan="2026-10-25").status_code, 200)
        self.assertEqual(self.add(name="Следом", d_from="2026-10-21", plan="2026-10-25").status_code, 200)
        self.assertEqual(self.add(sa="SKU-2", d_from="2026-10-20", plan="2026-10-19").status_code, 400)
        self.assertEqual(self.add(name="   ").status_code, 422)

    def test_close_edit_cancel_journal_and_history(self):
        pid = self.add().json()["id"]
        close = lambda uid, d: self.call(uid, "POST", f"/api/rnp-promos/{pid}/close", json={"date_to_fact": d})
        self.assertEqual(close(104, "2026-10-09").status_code, 400)  # раньше входа
        self.assertEqual(close(104, "2026-10-21").status_code, 400)  # позже плана — это продление
        self.assertEqual(close(105, "2026-10-15").status_code, 403)  # чужая акция без права «чужую»
        r = close(104, "2026-10-15")
        self.assertEqual((r.status_code, r.json()["status"], r.json()["date_to"], r.json()["days"]),
                         (200, "closed", "2026-10-15", 6))
        # После досрочного выхода место свободно: новая акция с 16.10.
        self.assertEqual(self.add(name="После", d_from="2026-10-16", plan="2026-10-18").status_code, 200)
        patch = lambda body: self.call(104, "PATCH", f"/api/rnp-promos/{pid}", json=body)
        self.assertEqual(patch({"date_to_plan": "2026-10-14"}).status_code, 400)  # не включает 15.10
        r = patch({"name": "Осенний драйв+"})
        self.assertEqual((r.status_code, r.json()["name"]), (200, "Осенний драйв+"))
        # Отменить: чужую — только с правом «чужую» (у администратора все права).
        self.assertEqual(self.call(105, "POST", f"/api/rnp-promos/{pid}/cancel").status_code, 403)
        r = self.call(101, "POST", f"/api/rnp-promos/{pid}/cancel")
        self.assertEqual((r.status_code, r.json()["status"], r.json()["cancelled_by_name"]), (200, "cancelled", "Админ"))
        self.assertEqual([p[1] for p in self.listed("wb")], ["После"])  # отменённой в РНП нет
        # Журнал: отменённая остаётся; права — РОП, Маркетинг, менеджеры, администратор.
        self.assertEqual(self.call(103, "GET", "/api/rnp-promos/journal").status_code, 403)
        r = self.call(102, "GET", "/api/rnp-promos/journal?marketplace=wb")
        self.assertEqual({p["name"]: p["status"] for p in r.json()["promos"]}["Осенний драйв+"], "cancelled")
        self.assertEqual(len(r.json()["promos"]), 2)
        ev = self.call(102, "GET", f"/api/rnp-promos/{pid}/events").json()["events"]
        self.assertEqual([e["action"] for e in ev], ["create", "close", "edit", "cancel"])
        self.assertEqual(ev[2]["changes"], {"name": ["Осенний драйв", "Осенний драйв+"]})
        self.assertEqual(ev[3]["user_name"], "Админ")

    def test_view_only_role_cannot_add(self):
        self.assertEqual(self.add(uid=103).status_code, 403)
        self.assertEqual(self.call(103, "GET", "/api/rnp-promos?marketplace=wb&date_from=2026-10-01"
                                             "&date_to=2026-10-31").status_code, 200)

    def test_status(self):
        r = {"cancelled_at": None, "date_to_fact": None, "date_from": D(2026, 10, 10), "date_to_plan": D(2026, 10, 20)}
        st = lambda today, **kw: rnp_promos_router._status({**r, **kw}, today)
        self.assertEqual(st(D(2026, 10, 9)), "planned")
        self.assertEqual(st(D(2026, 10, 10)), "active")
        self.assertEqual(st(D(2026, 10, 20)), "active")
        self.assertEqual(st(D(2026, 10, 21)), "finished")
        self.assertEqual(st(D(2026, 10, 12), date_to_fact=D(2026, 10, 12)), "closed")
        self.assertEqual(st(D(2026, 10, 12), cancelled_at=datetime.datetime.now()), "cancelled")


if __name__ == "__main__":
    unittest.main()
