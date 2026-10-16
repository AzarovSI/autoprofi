"""Загрузчик MPPROFIT Ozon: индекс заголовков, сверка с ИТОГО, общие удержания (без БД)."""
import os
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl

from app import loader_core as lc

# Раскладка как в отчёте MPPROFIT Ozon: секция «Налог» стоит над «Налоговой базой».
IDS = ["Артикул продавца", "SKU", "Наименование"]
METRICS = [  # (секция строки 4 или None, метрика строки 5)
    ("Продажи", "Заказы шт."), (None, "Заказы, руб."), (None, "Выкупы шт."), (None, "Выкупы, руб."),
    (None, "Возвраты руб."), (None, "Продажи руб."), (None, "Продажи шт."),
    ("Себестоимость", "Себестоимость"), ("Комиссии", "Комиссия"),
    ("Логистика", "Прямая логистика"), (None, "Выдача товара"), (None, "Вся логистика"),
    ("Удержания OZON", "Хранение"), (None, "Временное размещение"),
    ("Все удержания OZON", "Все удержания OZON"), ("Налог", "Налоговая база"), (None, "Налог"),
    ("Финансы", "Выплата на р/с"), (None, "Прибыль"), (None, "Марж-ность"),
]


def item(art, sku, revenue, commission, logistics, temp=0.0):
    holds = commission + logistics + temp
    return {"Артикул продавца": art, "SKU": sku, "Наименование": "Товар " + art,
            "Заказы шт.": 5, "Заказы, руб.": revenue * 1.3, "Выкупы шт.": 4, "Выкупы, руб.": revenue, "Возвраты руб.": 0,
            "Продажи руб.": revenue, "Продажи шт.": 4, "Себестоимость": 100,
            "Комиссия": commission, "Прямая логистика": logistics, "Вся логистика": logistics,
            "Временное размещение": temp, "Все удержания OZON": holds,
            "Выплата на р/с": revenue - holds, "Прибыль": revenue - holds - 100}


def holds_row(holds, extra_revenue=0.0):
    # Как в 12.2024: выплата строки удержаний учитывает выручку, которой нет в строках.
    return {"Артикул продавца": "Удержания OZON", "SKU": "Удержания OZON",
            "Хранение": holds, "Все удержания OZON": holds,
            "Выплата на р/с": extra_revenue - holds, "Прибыль": extra_revenue - holds}


def make_ws(rows, total_override=None, drop=()):
    ws = openpyxl.Workbook().active
    ws.cell(2, 3, "Период: 2026-09-01 - 2026-09-30")
    cols = IDS + [m for _, m in METRICS if m not in drop]
    sections = dict((m, s) for s, m in METRICS)
    for c, name in enumerate(cols, 1):
        if name in IDS:
            ws.cell(4, c, name)
        else:
            if sections.get(name):
                ws.cell(4, c, sections[name])
            ws.cell(5, c, name)
    for r, row in enumerate(rows, 7):
        for c, name in enumerate(cols, 1):
            if name in row:
                ws.cell(r, c, row[name])
    for c, name in enumerate(cols, 1):
        if name not in IDS:
            ws.cell(6, c, sum(row.get(name, 0) or 0 for row in rows))
    for name, v in (total_override or {}).items():
        ws.cell(6, cols.index(name) + 1, v)
    return ws


def idx_of(ws):
    idx = lc.build_colidx(ws, metric_row_first=True)
    lc.apply_oz_aliases(idx)
    return idx


ROWS = [item("A-1", "101", 1000, 300, 50, temp=12), item("B-2", "202", 2000, 500, 80), holds_row(40)]


class OzCheckTest(unittest.TestCase):
    def check(self, rows=ROWS, **kw):
        ws = make_ws(rows, **kw)
        return lc.check_oz_report(ws, idx_of(ws))

    def test_tax_reads_tax_column_not_section(self):
        ws = make_ws(ROWS)
        idx = idx_of(ws)
        self.assertEqual(ws.cell(5, idx["Налог"]).value, "Налог")

    def test_consistent_report_passes(self):
        self.assertEqual(self.check(), ([], [], {}))

    def test_strict_column_mismatch_rejected(self):
        errors, _, _ = self.check(total_override={"Комиссия": 900})
        self.assertTrue(any("«Комиссия»" in e for e in errors))

    def test_cancels_total_mismatch_only_warns(self):
        # Нед. 36/2026: ИТОГО невыкупов MPPROFIT 871 при Σ строк 877.
        rows = [dict(ROWS[0], **{"Невыкупы/отмены шт.": 4}), ROWS[1], ROWS[2]]
        ws = make_ws(rows)
        ws.cell(5, ws.max_column + 1, "Невыкупы/отмены шт.")
        c = ws.max_column
        ws.cell(7, c, 4); ws.cell(6, c, 2)
        errors, warnings, _ = lc.check_oz_report(ws, idx_of(ws))
        self.assertEqual(errors, [])
        self.assertIn("«Невыкупы/отмены шт.»: в ИТОГО 2,00, сумма строк 4,00", warnings[0])

    def test_orders_rub_total_mismatch_only_warns(self):
        errors, warnings, _ = self.check(total_override={"Заказы, руб.": 1})
        self.assertEqual(errors, [])
        self.assertIn("«Заказы, руб.»: в ИТОГО 1,00, сумма строк 3 900,00", warnings[0])

    def test_old_report_without_orders_rub_warns(self):
        # Отчёты Ozon до ~07.2026: колонки нет — грузим, на сайте остаётся оценка.
        ws = make_ws(ROWS, drop=("Заказы, руб.",))
        errors, warnings, _ = lc.check_oz_report(ws, idx_of(ws))
        self.assertEqual(errors, [])
        self.assertIn("«Заказы, руб.»", warnings[0])

    def test_handout_sign_quirk_ignored(self):
        rows = [dict(ROWS[0], **{"Выдача товара": -90.55}), ROWS[1], ROWS[2]]
        self.assertEqual(self.check(rows, total_override={"Выдача товара": 90.55})[0], [])

    def test_revenue_only_in_total_goes_to_holds_row(self):
        # 12.2024: ИТОГО выручки больше Σ строк, выплата строки удержаний это уже учла.
        rows = [ROWS[0], ROWS[1], holds_row(40, extra_revenue=500)]
        total = 3000 + 500
        errors, _, deltas = self.check(rows, total_override={"Продажи руб.": total, "Выкупы, руб.": total})
        self.assertEqual(errors, [])
        self.assertEqual(deltas, {"revenue": 500, "sales_gross": 500})

    def test_inconsistent_holds_row_rejected(self):
        errors, _, _ = self.check(total_override={"Продажи руб.": 3500, "Выкупы, руб.": 3500})
        self.assertTrue(any("«Удержания OZON» не сходится" in e for e in errors))

    def test_orphan_row_with_sku_counted(self):
        orphan = {"SKU": "202", "Выкупы, руб.": 2517, "Возвраты руб.": -2517,
                  "Прямая логистика": 189.07, "Вся логистика": 189.07,
                  "Все удержания OZON": 189.07, "Выплата на р/с": -189.07, "Прибыль": -189.07}
        self.assertEqual(self.check(ROWS + [orphan])[0], [])

    def test_merge_orphan_by_sku(self):
        cols = ["row_type", "seller_article", "sku_ozon"] + lc.OZ_MERGE_COLS
        row = lambda rt, art, sku, **v: tuple({"row_type": rt, "seller_article": art, "sku_ozon": sku, **v}.get(c) for c in cols)
        rows = [row("Товар", "B-2", "202", logistics_total=80, holds_total=580, payout=1420)]
        orphan = {"sku_ozon": "202", "logistics_total": 189.07, "holds_total": 189.07, "payout": -189.07}
        out, warnings = lc.merge_orphans(rows, [(49, orphan)], cols, lambda: None, "ozon")
        get = lambda col: out[0][cols.index(col)]
        self.assertEqual((get("logistics_total"), round(get("payout"), 2)), (269.07, 1230.93))
        self.assertIn("sku 202", warnings[0])
        self.assertIn("к товару B-2", warnings[0])


class OzLoadOrdersRubTest(unittest.TestCase):
    def load(self, drop=()):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            ws = make_ws(ROWS, drop=drop)
            ws.parent.save(path)
            conn = MagicMock()
            cur = conn.cursor.return_value.__enter__.return_value
            res = lc.load_report(conn, path, "ozon")
        finally:
            os.remove(path)
        ins = [c for c in cur.execute.call_args_list if "INSERT INTO fact_" in c.args[0]]
        self.assertTrue(ins, res)
        sql, flat = ins[0].args
        cols = sql[sql.index("(") + 1:sql.index(")")].split(",")
        rows = [flat[i:i + len(cols)] for i in range(0, len(flat), len(cols))]
        return {r[cols.index("seller_article")]: r[cols.index("orders_rub")] for r in rows}

    def test_orders_rub_written_from_report(self):
        got = self.load()
        self.assertEqual((got["A-1"], got["B-2"]), (1300, 2600))

    def test_old_report_leaves_orders_rub_empty(self):
        got = self.load(drop=("Заказы, руб.",))
        self.assertEqual((got["A-1"], got["B-2"]), (None, None))


if __name__ == "__main__":
    unittest.main()
