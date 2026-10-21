# -*- coding: utf-8 -*-
"""«Цены → Ценовая модель»: вводимые параметры с историей (таблица price_model_hist) —
цена категории, плановый ДРР и плановая маржа; пороги сигнала корректировки (price_model_rules).

История — одна запись на артикул + площадку + параметр + день (по Москве), правка в тот же
день перезаписывает значение дня. Загружаемая цена считается на фронте
(static/js/views.js, _pmCalcUpload). Правила — docs/business_rules.md.
"""
import math
from pathlib import Path

from . import db

# По порядку: 2026_10_21_price_category удаляет лист 13.10 и создаёт price_category_hist,
# 2026_10_21_price_model_params переносит её строки в общую price_model_hist,
# 2026_10_21_price_model_rules — пороги сигнала (⚙).
# DROP требует владельца таблиц: в TimeWeb применять под gen_user до выкладки.
MIGRATIONS = ["2026_10_21_price_category", "2026_10_21_price_model_params", "2026_10_21_price_model_rules"]

# Параметр → (единицы, нижняя граница, верхняя граница, функция-право на правку). Проценты — доли.
PARAMS = {
    "category_price": ("rub", 1, None, "prices.model.edit"),
    "drr_pct": ("pct", 0, 1, "prices.model.drr"),
    "margin_pct": ("pct", -0.5, 1, "prices.model.margin"),
}

# Порог сигнала корректировки по площадке (доля) — как в миграции (кнопка «Сбросить к исходным»).
DEFAULT_RULES = {"Wildberries": 0.05, "Ozon": 0.05}
RULE_MIN, RULE_MAX = 0.001, 0.5   # 0,1% … 50%, шаг 0,1 процента


def migrate():
    """Миграции по порядку: каждая один раз, атомарно, безопасно для нескольких воркеров."""
    with db.transaction() as tx:
        tx.execute("SELECT pg_advisory_xact_lock(13102026, 0)")
        tx.execute("""CREATE TABLE IF NOT EXISTS app_schema_migrations (
            name text PRIMARY KEY, applied_at timestamptz NOT NULL DEFAULT now())""")
        for name in MIGRATIONS:
            if tx.query_one("SELECT name FROM app_schema_migrations WHERE name=%s", (name,)):
                continue
            path = Path(__file__).resolve().parent.parent / "migrations" / (name + ".sql")
            tx.conn.execute(path.read_text(encoding="utf-8"))
            tx.execute("INSERT INTO app_schema_migrations(name) VALUES (%s)", (name,))


def load_rules():
    """{площадка: порог-доля}; чего нет в таблице — исходное значение."""
    rules = dict(DEFAULT_RULES)
    for r in db.query_all("SELECT marketplace, threshold FROM price_model_rules"):
        if r["marketplace"] in rules:
            rules[r["marketplace"]] = float(r["threshold"])
    return rules


def save_rules(rules, user_id):
    """Записать пороги {площадка: доля} одной транзакцией; ValueError с понятным текстом."""
    rows = []
    for mp, value in (rules or {}).items():
        if mp not in DEFAULT_RULES:
            raise ValueError(f"Неизвестная площадка «{mp}»")
        try:
            v = float(value)
        except (TypeError, ValueError):
            raise ValueError(f"{mp}: нужно число")
        if not math.isfinite(v) or v < RULE_MIN or v > RULE_MAX or abs(v * 1000 - round(v * 1000)) > 1e-6:
            raise ValueError(f"{mp}: порог — от {RULE_MIN * 100:g} до {RULE_MAX * 100:g}%, "
                             "не больше одного знака после запятой")
        rows.append((mp, round(v * 1000) / 1000))
    with db.transaction() as tx:
        for mp, v in rows:
            tx.execute(
                """INSERT INTO price_model_rules (marketplace, threshold, updated_at, updated_by_user_id)
                   VALUES (%s, %s, now(), %s)
                   ON CONFLICT (marketplace) DO UPDATE
                   SET threshold = EXCLUDED.threshold, updated_at = now(),
                       updated_by_user_id = EXCLUDED.updated_by_user_id""",
                (mp, v, user_id))
    return load_rules()


# ── Значения параметров: проверка, срез, запись с историей ──────────────────
def check_value(param, value):
    """Проверить и привести значение параметра (None — удалить). ValueError с текстом для человека.
    Цена категории — целое > 0; проценты — доли с шагом 0,1 процента в границах PARAMS."""
    if value is None:
        return None
    unit, lo, hi, _ = PARAMS[param]
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("нужно число")
    if unit == "rub":
        if value != int(value) or value < lo:
            raise ValueError("цена категории — целое число больше нуля")
        return int(value)
    if value < lo or value > hi or abs(value * 1000 - round(value * 1000)) > 1e-6:
        raise ValueError(f"от {lo * 100:g} до {hi * 100:g}%, не больше одного знака после запятой")
    return round(value * 1000) / 1000


def value_out(param, v):
    """Значение из БД для ответа: цена — int, проценты — float."""
    if v is None:
        return None
    return int(v) if PARAMS[param][0] == "rub" else float(v)


def load_slice(dt):
    """Срез на дату dt: по каждой тройке артикул + площадка + параметр — последняя запись
    с датой ≤ dt; удалённые (NULL) не отдаются. У цены категории — ещё СПП из ценового
    отчёта за день, с которого она действует (подсказка сигнала); оценочная не берётся."""
    rows = db.query_all(
        """SELECT s.*, CASE WHEN s.param = 'category_price' AND NOT COALESCE(m.spp_is_estimated, false)
                            THEN m.spp_pct END AS spp_at_date
             FROM (
               SELECT DISTINCT ON (h.seller_article, h.marketplace, h.param)
                      h.seller_article, h.marketplace, h.param, h.value, h.date,
                      COALESCE(u.display_name, u.email) AS updated_by
                 FROM price_model_hist h
                 LEFT JOIN app_users u ON u.id = h.updated_by_user_id
                WHERE h.date <= %s
                ORDER BY h.seller_article, h.marketplace, h.param, h.date DESC
           ) s
             LEFT JOIN mp_price_daily m
                    ON m.seller_article = s.seller_article AND m.marketplace = s.marketplace AND m.date = s.date
          WHERE s.value IS NOT NULL
          ORDER BY s.seller_article, s.marketplace, s.param""",
        (dt,))
    for r in rows:
        r["value"] = value_out(r["param"], r["value"])
        r["date"] = r["date"].isoformat()
        r["spp_at_date"] = float(r["spp_at_date"]) if r["spp_at_date"] is not None else None
    return rows


def _same(a, b):
    return (a is None and b is None) or (a is not None and b is not None and abs(float(a) - float(b)) < 1e-9)


def write_values(entries, user_id, today):
    """Записать значения на день today одной транзакцией, пачками.

    entries — [(артикул, площадка, параметр, значение|None)], значения уже проверены
    (check_value). Одна запись на тройку и день: правка дня перезаписывает значение дня.
    Если значение равно действовавшему до today — запись дня не нужна (удаляется, если была):
    история не пополняется повторами. Возвращает {"set", "cleared"} — сколько значений
    реально задано/изменено и сколько удалено (относительно действующего на сейчас)."""
    entries = list({(a, mp, p): (a, mp, p, v) for a, mp, p, v in entries}.values())
    if not entries:
        return {"set": 0, "cleared": 0}
    keys = ([e[0] for e in entries], [e[1] for e in entries], [e[2] for e in entries])
    with db.transaction() as tx:
        prev, cur = {}, {}
        for r in tx.query_all(
                """SELECT DISTINCT ON (seller_article, marketplace, param)
                          seller_article, marketplace, param, value
                     FROM price_model_hist
                    WHERE (seller_article, marketplace, param) IN (
                          SELECT unnest(%s::text[]), unnest(%s::text[]), unnest(%s::text[]))
                      AND date < %s
                    ORDER BY seller_article, marketplace, param, date DESC""", (*keys, today)):
            prev[(r["seller_article"], r["marketplace"], r["param"])] = r["value"]
        for r in tx.query_all(
                """SELECT seller_article, marketplace, param, value
                     FROM price_model_hist
                    WHERE (seller_article, marketplace, param) IN (
                          SELECT unnest(%s::text[]), unnest(%s::text[]), unnest(%s::text[]))
                      AND date = %s""", (*keys, today)):
            cur[(r["seller_article"], r["marketplace"], r["param"])] = r["value"]
        drop, upsert, n_set, n_cleared = [], [], 0, 0
        for a, mp, p, v in entries:
            k = (a, mp, p)
            now_val = cur[k] if k in cur else prev.get(k)
            if not _same(now_val, v):
                if v is None:
                    n_cleared += 1
                else:
                    n_set += 1
            if _same(prev.get(k), v):
                if k in cur:
                    drop.append(k)
            else:
                upsert.append((a, mp, p, today, v, user_id))
        if drop:
            tx.execute(
                """DELETE FROM price_model_hist
                    WHERE (seller_article, marketplace, param) IN (
                          SELECT unnest(%s::text[]), unnest(%s::text[]), unnest(%s::text[]))
                      AND date = %s""",
                ([k[0] for k in drop], [k[1] for k in drop], [k[2] for k in drop], today))
        if upsert:
            tx.execute_values(
                """INSERT INTO price_model_hist (seller_article, marketplace, param, date, value,
                                                 updated_at, updated_by_user_id)
                   VALUES %s
                   ON CONFLICT (seller_article, marketplace, param, date) DO UPDATE SET
                        value = EXCLUDED.value,
                        updated_at = now(),
                        updated_by_user_id = EXCLUDED.updated_by_user_id""",
                upsert, template="(%s,%s,%s,%s,%s,now(),%s)")
    return {"set": n_set, "cleared": n_cleared}


# ── Excel: шаблон и загрузка (этап 2 плана владельца 21.10.2026) ───────────
# Колонки значений — площадка × параметр; при загрузке ищутся по шапке, справочные
# колонки (иерархия, название) не читаются. Пустая ячейка — удалить значение.
EXCEL_MPS = (("Wildberries", "WB"), ("Ozon", "Ozon"))
EXCEL_PARAMS = (("drr_pct", "ДРР"), ("margin_pct", "Маржа"), ("category_price", "Цена категории"))
EXCEL_REF = ("Группа", "Подгруппа", "Категория", "Артикул", "Название")


def excel_columns():
    """[(заголовок, площадка, параметр)] — колонки значений шаблона по порядку."""
    return [(f"{plabel} {mlabel}", mp, param) for mp, mlabel in EXCEL_MPS for param, plabel in EXCEL_PARAMS]


def export_workbook(items, values):
    """Шаблон: строки items (как в отчёте), значения — {(артикул, площадка, параметр): value}."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .price_export import text_cell

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ценовая модель"
    cols = excel_columns()
    header = list(EXCEL_REF) + [c[0] for c in cols]
    ws.append(header)
    fills = {"Wildberries": PatternFill("solid", fgColor="F3E2EF"), "Ozon": PatternFill("solid", fgColor="E3EBFB")}
    for ci, title in enumerate(header, 1):
        cell = ws.cell(row=1, column=ci)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ci > len(EXCEL_REF):
            cell.fill = fills[cols[ci - len(EXCEL_REF) - 1][1]]
    for it in items:
        art = it.get("seller_article")
        row = [it.get("l1"), it.get("l2"), it.get("l3"), art, it.get("sample_name")]
        row += [values.get((art, mp, param)) for _, mp, param in cols]
        ws.append(row)
        r = ws.max_row
        for ci in range(1, len(EXCEL_REF) + 1):
            text_cell(ws.cell(r, ci), row[ci - 1])
        for ci, (_, _, param) in enumerate(cols, len(EXCEL_REF) + 1):
            ws.cell(r, ci).number_format = "0.0%" if PARAMS[param][0] == "pct" else "#,##0"
    for ci, w in enumerate((22, 22, 22, 18, 40) + (12, 12, 14) * len(EXCEL_MPS), 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _excel_pct(cell):
    """Процент из ячейки → доля. Ячейка в процентном формате (12,5%) — уже доля;
    число без процентного формата и текст — проценты (12,5 → 0,125). Пусто — None."""
    v = cell.value
    if v is None or (isinstance(v, str) and v.strip() in ("", "-", "—")):
        return None
    if isinstance(v, bool):
        raise ValueError("нужно число")
    if isinstance(v, (int, float)):
        return float(v) if "%" in (getattr(cell, "number_format", "") or "") else float(v) / 100
    s = str(v).replace(" ", "").replace(" ", "").replace("%", "").replace(",", ".")
    try:
        return float(s) / 100
    except ValueError:
        raise ValueError("нужно число")


def _excel_rub(cell):
    v = cell.value
    if v is None or (isinstance(v, str) and v.strip() in ("", "-", "—")):
        return None
    if isinstance(v, bool):
        raise ValueError("нужно число")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        raise ValueError("нужно число")


def _excel_shown(cell):
    """Значение ячейки для сообщения об ошибке — как его видит сотрудник (150%, а не 1.5)."""
    v = None if cell is None else cell.value
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (getattr(cell, "number_format", "") or ""):
        return f"{v * 100:g}%".replace(".", ",")
    return str(v)


def _excel_header(title):
    """(площадка, параметр) по заголовку колонки или None."""
    s = str(title or "").strip().lower()
    if "ozon" in s or "озон" in s:
        mp = "Ozon"
    elif "wb" in s or "wildberries" in s or "вб" in s:
        mp = "Wildberries"
    else:
        return None
    if "дрр" in s:
        return mp, "drr_pct"
    if "марж" in s:
        return mp, "margin_pct"
    if "категор" in s:
        return mp, "category_price"
    return None


def parse_workbook(data, allowed):
    """Разобрать загруженный шаблон.

    allowed(param) — может ли пользователь менять параметр: колонки без права не читаются.
    Возвращает {"rows": [(строка, артикул, {(площадка, параметр): значение})], "skipped_columns": [...],
    "errors": [...]}; значения уже проверены check_value. ValueError — файл не разобрать целиком."""
    import io
    import openpyxl
    from . import util

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = list(wb.active.iter_rows())
    except Exception:
        raise ValueError("Не удалось прочитать Excel-файл")
    if not rows:
        raise ValueError("Файл пустой")
    art_col, cols, skipped = None, {}, []
    for ci, cell in enumerate(rows[0]):
        title = str(cell.value or "").strip()
        if art_col is None and title.lower() == "артикул":
            art_col = ci
            continue
        key = _excel_header(title)
        if key is None or key in cols.values():
            continue
        if allowed(key[1]):
            cols[ci] = key
        elif title not in skipped:
            skipped.append(title)
    if art_col is None:
        raise ValueError("В шапке файла нет колонки «Артикул»")
    if not cols:
        raise ValueError("В файле нет колонок, которые вы можете загружать"
                         + (f" (без права: {', '.join(skipped)})" if skipped else ""))
    out, errors, seen = [], [], {}
    for ri, row in enumerate(rows[1:], 2):
        cell = row[art_col] if art_col < len(row) else None
        art = util.canon_article(cell.value if cell is not None else None)
        if not art:
            continue
        if art in seen:
            errors.append(f"строка {ri}: артикул {art} уже был в строке {seen[art]}")
            continue
        seen[art] = ri
        vals = {}
        for ci, (mp, param) in cols.items():
            c = row[ci] if ci < len(row) else None
            title = rows[0][ci].value
            try:
                raw = None if c is None else (_excel_pct(c) if PARAMS[param][0] == "pct" else _excel_rub(c))
                vals[(mp, param)] = check_value(param, raw)
            except ValueError as e:
                errors.append(f"строка {ri}, «{title}»: «{_excel_shown(c)}» — {e}")
        out.append((ri, art, vals))
    return {"rows": out, "skipped_columns": skipped, "errors": errors}
