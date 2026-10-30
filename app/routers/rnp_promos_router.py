# -*- coding: utf-8 -*-
"""Роутер акций товаров в «РНП заказы» (решение владельца 23.10.2026).

Акция = товар площадки (marketplace, seller_article) + название + вход (date_from) + плановый выход
(date_to_plan) + фактический выход при досрочном завершении (date_to_fact). Даты включительно.
У товара на площадке одна акция одновременно. Записи не удаляются: «отменить» помечает акцию
(cancelled_at), она пропадает из РНП, но остаётся в журнале. Каждое действие пишется в
rnp_promo_events — для журнала и статистики. Правила — docs/business_rules.md, «Акции в РНП».
Открыть / изменить / завершить / отменить / загрузить = функция действия + право на площадку
(rnp.promos.mp.wb | ozon | ya, решение владельца 30.10.2026); смотреть можно все площадки.

Эндпоинты:
  GET    /api/rnp-promos?marketplace=&date_from=&date_to=  — акции площадки для РНП
  POST   /api/rnp-promos                                    — открыть акцию
  PATCH  /api/rnp-promos/{id}                               — изменить название и даты
  POST   /api/rnp-promos/{id}/close                         — завершить раньше
  POST   /api/rnp-promos/{id}/cancel                        — отменить запись
  GET    /api/rnp-promos/journal?...                        — журнал всех акций
  GET    /api/rnp-promos/{id}/events                        — история изменений акции
  GET    /api/rnp-promos/template                           — пустой Excel-шаблон «одна акция»
  POST   /api/rnp-promos/import                             — загрузить заполненный шаблон (только добавление)
"""
import datetime
import io
import json
from datetime import date as date_cls
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, field_validator

from .. import auth, db, util
from .rnp_sales_router import MP_MAP

router = APIRouter(prefix="/api/rnp-promos", tags=["rnp-promos"])

MIGRATION = "2026_10_23_rnp_promos"
IMPORT_MIGRATION = "2026_10_30_rnp_promos_import"
MIGRATIONS = [MIGRATION, IMPORT_MIGRATION]
MAX_NAME = 200
_MSK = ZoneInfo("Europe/Moscow")


def migrate():
    """Таблицы акций и права ролей: каждая миграция один раз, по порядку, атомарно, безопасно для
    нескольких воркеров. Запускается после permissions.migrate() — права выдаются по таблицам ролей."""
    with db.transaction() as tx:
        tx.execute("SELECT pg_advisory_xact_lock(23102026, 1)")
        tx.execute("""CREATE TABLE IF NOT EXISTS app_schema_migrations (
            name text PRIMARY KEY, applied_at timestamptz NOT NULL DEFAULT now())""")
        for name in MIGRATIONS:
            if tx.query_one("SELECT name FROM app_schema_migrations WHERE name=%s", (name,)):
                continue
            path = Path(__file__).resolve().parents[2] / "migrations" / (name + ".sql")
            tx.conn.execute(path.read_text(encoding="utf-8"))
            tx.execute("INSERT INTO app_schema_migrations(name) VALUES (%s)", (name,))


def _today():
    return datetime.datetime.now(_MSK).date()


def _mp(value):
    """Площадка из параметра (ozon | wb | yandex) → имя как в БД; неизвестная — ошибка."""
    mp = MP_MAP.get((value or "").strip().lower())
    if not mp:
        raise ValueError("Неизвестная площадка")
    return mp


_MP_RIGHT = {"Wildberries": "rnp.promos.mp.wb", "Ozon": "rnp.promos.mp.ozon", "Yandex": "rnp.promos.mp.ya"}


def _check_mp(user, mp):
    """403, если у роли нет права на акции этой площадки (смотреть акции можно все)."""
    auth.check(user, _MP_RIGHT[mp])


def _name(v):
    v = " ".join((v or "").split())
    if not v:
        raise ValueError("Укажите название акции")
    if len(v) > MAX_NAME:
        raise ValueError(f"Название длиннее {MAX_NAME} символов")
    return v


class PromoCreate(BaseModel):
    marketplace: str
    seller_article: str
    name: str
    date_from: date_cls
    date_to_plan: date_cls

    @field_validator("marketplace")
    @classmethod
    def _market(cls, v):
        return _mp(v)

    @field_validator("seller_article")
    @classmethod
    def _art(cls, v):
        v = (v or "").strip()
        if not v:
            raise ValueError("Не указан артикул")
        return v

    @field_validator("name")
    @classmethod
    def _nm(cls, v):
        return _name(v)


class PromoUpdate(BaseModel):
    name: Optional[str] = None
    date_from: Optional[date_cls] = None
    date_to_plan: Optional[date_cls] = None

    @field_validator("name")
    @classmethod
    def _nm(cls, v):
        return None if v is None else _name(v)


class PromoClose(BaseModel):
    date_to_fact: date_cls


def _bad(msg):
    raise HTTPException(status_code=400, detail=msg)


def _status(r, today=None):
    """planned — ещё не началась, active — идёт, finished — закончилась по плану,
    closed — завершена раньше, cancelled — отменена."""
    today = today or _today()
    if r["cancelled_at"]:
        return "cancelled"
    if r["date_to_fact"]:
        return "closed"
    if today < r["date_from"]:
        return "planned"
    if today > r["date_to_plan"]:
        return "finished"
    return "active"


def _iso(v):
    return v.isoformat() if v else None


def _row_to_dict(r, today=None):
    end = r["date_to_fact"] or r["date_to_plan"]
    return {
        "id": r["id"],
        "marketplace": r["marketplace"],
        "seller_article": r["seller_article"],
        "name": r["name"],
        "date_from": _iso(r["date_from"]),
        "date_to_plan": _iso(r["date_to_plan"]),
        "date_to_fact": _iso(r["date_to_fact"]),
        "date_to": _iso(end),
        "days": (end - r["date_from"]).days + 1,
        "status": _status(r, today),
        "created_by": r["created_by"],
        "created_by_name": r.get("created_by_name"),
        "created_at": _iso(r["created_at"]),
        "closed_by_name": r.get("closed_by_name"),
        "closed_at": _iso(r["closed_at"]),
        "cancelled_by_name": r.get("cancelled_by_name"),
        "cancelled_at": _iso(r["cancelled_at"]),
        "source_file": r.get("source_file"),
    }


_SELECT = """
    SELECT p.*, uc.display_name AS created_by_name, ux.display_name AS closed_by_name,
           un.display_name AS cancelled_by_name,
           (SELECT e.changes->>'file' FROM rnp_promo_events e
            WHERE e.promo_id = p.id AND e.action = 'create' LIMIT 1) AS source_file
    FROM rnp_promos p
    LEFT JOIN app_users uc ON uc.id = p.created_by
    LEFT JOIN app_users ux ON ux.id = p.closed_by
    LEFT JOIN app_users un ON un.id = p.cancelled_by
"""


def _get(promo_id, tx=None):
    q = (tx or db).query_one(_SELECT + " WHERE p.id = %s", (promo_id,))
    if not q:
        raise HTTPException(status_code=404, detail="Акция не найдена")
    return q


def _lock_item(tx, mp, sa):
    """Сериализует правки акций одного товара площадки (проверка «одна акция» без гонок).
    Общая блокировка площадки — чтобы не пересечься с загрузкой из Excel (_lock_mp)."""
    tx.execute("SELECT pg_advisory_xact_lock_shared(hashtext(%s))", (f"rnp_promo|{mp}",))
    tx.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", (f"rnp_promo|{mp}|{sa}",))


def _lock_mp(tx, mp):
    """Загрузка из Excel: исключительная блокировка площадки — ждёт ручные правки её акций."""
    tx.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", (f"rnp_promo|{mp}",))


def _check_overlap(tx, mp, sa, d_from, d_to, exclude_id=None):
    """У товара на площадке одна акция одновременно: периоды не пересекаются."""
    other = tx.query_one(
        "SELECT name, date_from, COALESCE(date_to_fact, date_to_plan) AS date_to FROM rnp_promos "
        "WHERE marketplace = %s AND seller_article = %s AND cancelled_at IS NULL AND id IS DISTINCT FROM %s "
        "AND date_from <= %s AND COALESCE(date_to_fact, date_to_plan) >= %s "
        "ORDER BY date_from LIMIT 1",
        (mp, sa, exclude_id, d_to, d_from),
    )
    if other:
        _bad(f"У товара уже есть акция «{other['name']}» "
             f"{other['date_from']:%d.%m.%Y} — {other['date_to']:%d.%m.%Y}. Две акции одновременно быть не может.")


def _event(tx, promo_id, action, changes, user):
    tx.execute("INSERT INTO rnp_promo_events (promo_id, action, changes, user_id) VALUES (%s, %s, %s::jsonb, %s)",
               (promo_id, action, json.dumps(changes, ensure_ascii=False, default=str), user["id"]))


def _assert_can_modify(promo, user):
    """Своя акция — функция «свою», чужая — «чужую» (как комментарии)."""
    if promo["created_by"] == user["id"]:
        auth.check(user, "rnp.promos.edit_own")
    else:
        auth.check(user, "rnp.promos.edit_any")


@router.get("")
def list_promos(
    marketplace: str = Query(..., description="площадка РНП: ozon | wb | yandex"),
    date_from: date_cls = Query(...),
    date_to: date_cls = Query(...),
    user=Depends(auth.require("rnp.promos.view")),
):
    """Неотменённые акции площадки, которые пересекают видимый период или идут сегодня
    (ракета у товара не зависит от выбранного периода)."""
    try:
        mp = _mp(marketplace)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    today = _today()
    rows = db.query_all(
        _SELECT + " WHERE p.marketplace = %s AND p.cancelled_at IS NULL "
        "AND ((p.date_from <= %s AND COALESCE(p.date_to_fact, p.date_to_plan) >= %s) "
        "  OR (p.date_from <= %s AND COALESCE(p.date_to_fact, p.date_to_plan) >= %s)) "
        "ORDER BY p.seller_article, p.date_from",
        (mp, date_to, date_from, today, today),
    )
    return {"today": today.isoformat(), "promos": [_row_to_dict(r, today) for r in rows]}


@router.post("")
def create_promo(payload: PromoCreate, user=Depends(auth.require("rnp.promos.add"))):
    _check_mp(user, payload.marketplace)
    if payload.date_to_plan < payload.date_from:
        _bad("Плановый выход раньше входа в акцию")
    with db.transaction() as tx:
        _lock_item(tx, payload.marketplace, payload.seller_article)
        _check_overlap(tx, payload.marketplace, payload.seller_article, payload.date_from, payload.date_to_plan)
        new_id = tx.query_one(
            "INSERT INTO rnp_promos (marketplace, seller_article, name, date_from, date_to_plan, created_by) "
            "VALUES (%s, %s, %s, %s, %s, %s) RETURNING id",
            (payload.marketplace, payload.seller_article, payload.name, payload.date_from,
             payload.date_to_plan, user["id"]),
        )["id"]
        _event(tx, new_id, "create", {"name": payload.name, "date_from": payload.date_from,
                                      "date_to_plan": payload.date_to_plan}, user)
        row = _get(new_id, tx)
    return _row_to_dict(row)


@router.patch("/{promo_id}")
def update_promo(promo_id: int, payload: PromoUpdate, user=Depends(auth.get_current_user)):
    """Изменить название, вход и плановый выход. Дата досрочного выхода остаётся, если
    по-прежнему внутри нового периода."""
    with db.transaction() as tx:
        cur = _get(promo_id, tx)
        _assert_can_modify(cur, user)
        _check_mp(user, cur["marketplace"])
        if cur["cancelled_at"]:
            _bad("Акция отменена — изменить нельзя")
        new = {"name": payload.name or cur["name"],
               "date_from": payload.date_from or cur["date_from"],
               "date_to_plan": payload.date_to_plan or cur["date_to_plan"]}
        if new["date_to_plan"] < new["date_from"]:
            _bad("Плановый выход раньше входа в акцию")
        fact = cur["date_to_fact"]
        if fact and not (new["date_from"] <= fact <= new["date_to_plan"]):
            _bad(f"Акция завершена {fact:%d.%m.%Y} — новые даты должны включать этот день")
        changes = {k: [cur[k], v] for k, v in new.items() if cur[k] != v}
        if not changes:
            return _row_to_dict(cur)
        _lock_item(tx, cur["marketplace"], cur["seller_article"])
        _check_overlap(tx, cur["marketplace"], cur["seller_article"], new["date_from"],
                       fact or new["date_to_plan"], exclude_id=promo_id)
        tx.execute("UPDATE rnp_promos SET name = %s, date_from = %s, date_to_plan = %s, updated_at = now() "
                   "WHERE id = %s", (new["name"], new["date_from"], new["date_to_plan"], promo_id))
        _event(tx, promo_id, "edit", changes, user)
        row = _get(promo_id, tx)
    return _row_to_dict(row)


@router.post("/{promo_id}/close")
def close_promo(promo_id: int, payload: PromoClose, user=Depends(auth.get_current_user)):
    """Завершить раньше: последний день акции = date_to_fact (не позже планового выхода)."""
    with db.transaction() as tx:
        cur = _get(promo_id, tx)
        _assert_can_modify(cur, user)
        _check_mp(user, cur["marketplace"])
        if cur["cancelled_at"]:
            _bad("Акция отменена")
        d = payload.date_to_fact
        if d < cur["date_from"]:
            _bad(f"Акция начинается {cur['date_from']:%d.%m.%Y} — завершить раньше входа нельзя")
        if d > cur["date_to_plan"]:
            _bad(f"Плановый выход — {cur['date_to_plan']:%d.%m.%Y}. Чтобы продлить акцию, измените плановый выход")
        tx.execute("UPDATE rnp_promos SET date_to_fact = %s, closed_by = %s, closed_at = now(), updated_at = now() "
                   "WHERE id = %s", (d, user["id"], promo_id))
        _event(tx, promo_id, "close", {"date_to_fact": [cur["date_to_fact"], d]}, user)
        row = _get(promo_id, tx)
    return _row_to_dict(row)


@router.post("/{promo_id}/cancel")
def cancel_promo(promo_id: int, user=Depends(auth.get_current_user)):
    """Отменить ошибочную запись: из РНП пропадает, в журнале остаётся с пометкой."""
    with db.transaction() as tx:
        cur = _get(promo_id, tx)
        _assert_can_modify(cur, user)
        _check_mp(user, cur["marketplace"])
        if cur["cancelled_at"]:
            return _row_to_dict(cur)
        tx.execute("UPDATE rnp_promos SET cancelled_by = %s, cancelled_at = now(), updated_at = now() "
                   "WHERE id = %s", (user["id"], promo_id))
        _event(tx, promo_id, "cancel", {}, user)
        row = _get(promo_id, tx)
    return _row_to_dict(row)


@router.get("/journal")
def journal(
    marketplace: Optional[str] = Query(None, description="ozon | wb | yandex; пусто — все"),
    date_from: Optional[date_cls] = Query(None, description="акции, пересекающие период"),
    date_to: Optional[date_cls] = Query(None),
    user=Depends(auth.require("rnp.promos.journal")),
):
    """Все акции (и отменённые) с авторами — для вкладки «Журнал акций». Фильтры по статусу и
    поиску — на клиенте."""
    where, params = [], []
    if marketplace:
        try:
            where.append("p.marketplace = %s")
            params.append(_mp(marketplace))
        except ValueError as e:
            raise HTTPException(status_code=422, detail=str(e))
    if date_to:
        where.append("p.date_from <= %s")
        params.append(date_to)
    if date_from:
        where.append("COALESCE(p.date_to_fact, p.date_to_plan) >= %s")
        params.append(date_from)
    rows = db.query_all(_SELECT + (" WHERE " + " AND ".join(where) if where else "")
                        + " ORDER BY p.date_from DESC, p.id DESC", params)
    today = _today()
    return {"today": today.isoformat(), "promos": [_row_to_dict(r, today) for r in rows]}


@router.get("/{promo_id}/events")
def events(promo_id: int, user=Depends(auth.require("rnp.promos.journal", "rnp.promos.view"))):
    """История акции: кто, когда и что менял."""
    _get(promo_id)
    rows = db.query_all(
        "SELECT e.action, e.changes, e.at, e.user_id, u.display_name AS user_name "
        "FROM rnp_promo_events e LEFT JOIN app_users u ON u.id = e.user_id "
        "WHERE e.promo_id = %s ORDER BY e.at, e.id", (promo_id,))
    return {"events": [{"action": r["action"], "changes": r["changes"], "at": _iso(r["at"]),
                        "user_name": r["user_name"]} for r in rows]}


# ===== Загрузка акции из Excel (решения владельца 30.10.2026) =====
# Шаблон «один файл = одна акция»: вверху площадка, название, вход и плановый выход, ниже — столбец
# артикулов. Только добавление: правильные строки записываются, ошибочные возвращаются списком с
# причиной (нет на площадке, уже в другой акции, повтор). Ошибка в шапке — не записывается ничего.
# Изменить, завершить и отменить акции — вручную, как раньше.

TPL_TITLE = "Товары в акцию — РНП ТД АВТОПРОФИ"
TPL_MP_NAMES = {"Wildberries": "Wildberries", "Ozon": "OZON", "Yandex": "Yandex"}
# Подпись в столбце A → поле шапки. Ищется по тексту, а не по номеру строки.
TPL_FIELDS = (("площадка", "mp", "Площадка"), ("название акции", "name", "Название акции"),
              ("вход в акцию", "date_from", "Вход в акцию"), ("плановый выход", "date_to_plan", "Плановый выход"))
TPL_ARTICLES = "артикул продавца"
TPL_FIRST_ROW = 10        # первая строка артикулов в выгружаемом шаблоне
MAX_IMPORT_ROWS = 5000
_MP_IMPORT = {"wildberries": "Wildberries", "wb": "Wildberries", "вб": "Wildberries", "вайлдберриз": "Wildberries",
              "ozon": "Ozon", "озон": "Ozon",
              "yandex": "Yandex", "яндекс": "Yandex", "яндекс маркет": "Yandex", "яндекс.маркет": "Yandex",
              "ya": "Yandex", "ям": "Yandex"}


@router.get("/template")
def template(user=Depends(auth.require("rnp.promos.import"))):
    """Пустой шаблон «одна акция»: шапка без товаров (товаров в акции много меньше, чем в ассортименте)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Акция"
    hint = Font(color="7A838C", italic=True)
    ws["A1"] = TPL_TITLE
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Заполните жёлтые ячейки и столбец артикулов. Один файл — одна акция одной площадки."
    ws["A2"].font = hint
    fill = PatternFill("solid", fgColor="FFF8D6")
    hints = ("список: " + " / ".join(TPL_MP_NAMES.values()), "как в кабинете площадки",
             "первый день акции, ДД.ММ.ГГГГ", "последний день акции, ДД.ММ.ГГГГ")
    for i, ((_, field, label), h) in enumerate(zip(TPL_FIELDS, hints)):
        r = 4 + i
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2).fill = fill
        if field in ("date_from", "date_to_plan"):
            ws.cell(r, 2).number_format = "DD.MM.YYYY"
        ws.cell(r, 3, h).font = hint
    dv = DataValidation(type="list", formula1='"' + ",".join(TPL_MP_NAMES.values()) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    head = ws.cell(TPL_FIRST_ROW - 1, 1, "Артикул продавца")
    head.font = Font(bold=True)
    head.fill = PatternFill("solid", fgColor="E7EEF7")
    head.alignment = Alignment(vertical="center")
    for r in range(TPL_FIRST_ROW, TPL_FIRST_ROW + 1000):
        ws.cell(r, 1).number_format = "@"   # артикул — текст: «00123» не станет числом 123
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 36
    ws.freeze_panes = f"A{TPL_FIRST_ROW}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = "Акция — шаблон загрузки.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}",
                 "Cache-Control": "no-cache, no-store, must-revalidate, private",
                 "Pragma": "no-cache", "Expires": "0"},
    )


def _label(v):
    return " ".join(str(v or "").replace(":", " ").split()).lower()


def _cell_text(v):
    """Текст ячейки: число 12345.0 → «12345», дата не ожидается."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _cell_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, date_cls):
        return v
    s = _cell_text(v)
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_template(content):
    """Шапка и артикулы шаблона. Ошибки шапки → ValueError со списком (файл не загружается).
    Возвращает {mp, name, date_from, date_to_plan, rows: [(номер строки, текст артикула)]}."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Не удалось прочитать файл — нужен Excel (.xlsx), выгруженный кнопкой «Шаблон акции»")
    ws = wb.worksheets[0]
    head, where, rows, art_row = {}, {}, [], None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        a = row[0] if row else None
        if art_row is None:
            lbl = _label(a)
            if lbl == TPL_ARTICLES:
                art_row = i
                continue
            for key, field, _ in TPL_FIELDS:
                if lbl == key and field not in head:
                    head[field] = row[1] if len(row) > 1 else None
                    where[field] = f"B{i}"
            continue
        text = _cell_text(a)
        if text:
            rows.append((i, text))
            if len(rows) > MAX_IMPORT_ROWS:
                raise ValueError(f"В файле больше {MAX_IMPORT_ROWS} артикулов — разбейте его на несколько")
    wb.close()

    errors = []
    for key, field, label in TPL_FIELDS:
        if field not in head:
            errors.append(f"нет строки «{label}» — выгрузите шаблон заново")
    if art_row is None:
        errors.append("нет строки «Артикул продавца» — выгрузите шаблон заново")
    if errors:
        raise ValueError("; ".join(errors))

    mp_text = _cell_text(head["mp"])
    mp = _MP_IMPORT.get(mp_text.lower())
    if not mp_text:
        errors.append(f"не выбрана площадка (ячейка {where['mp']})")
    elif not mp:
        errors.append(f"неизвестная площадка «{mp_text}» (ячейка {where['mp']}) — выберите "
                      + ", ".join(TPL_MP_NAMES.values()))
    name = _cell_text(head["name"])
    if not name:
        errors.append(f"не указано название акции (ячейка {where['name']})")
    elif len(name) > MAX_NAME:
        errors.append(f"название акции длиннее {MAX_NAME} символов")
    dates = {}
    for field, label in (("date_from", "вход в акцию"), ("date_to_plan", "плановый выход")):
        raw = head[field]
        d = _cell_date(raw)
        if raw in (None, ""):
            errors.append(f"не указан {label} (ячейка {where[field]})")
        elif not d:
            errors.append(f"{label} «{_cell_text(raw)}» (ячейка {where[field]}) — не дата, нужна в виде ДД.ММ.ГГГГ")
        dates[field] = d
    if dates["date_from"] and dates["date_to_plan"] and dates["date_to_plan"] < dates["date_from"]:
        errors.append("плановый выход раньше входа в акцию")
    if not rows and not errors:
        errors.append("в столбце «Артикул продавца» нет ни одного артикула")
    if errors:
        raise ValueError("; ".join(errors))
    return {"mp": mp, "name": name, "date_from": dates["date_from"], "date_to_plan": dates["date_to_plan"],
            "rows": rows}


@router.post("/import")
def import_promo(file: UploadFile = File(...), user=Depends(auth.require("rnp.promos.import"))):
    """Загрузить заполненный шаблон: правильные артикулы — в акцию (как ручное «Начать акцию»),
    остальные — списком с причиной. Повторная загрузка того же файла товары не задваивает."""
    try:
        t = parse_template(file.file.read())
    except ValueError as e:
        raise HTTPException(status_code=422, detail=f"Файл не загружен: {e}")
    mp, name, d_from, d_to = t["mp"], t["name"], t["date_from"], t["date_to_plan"]
    _check_mp(user, mp)
    fname = (file.filename or "").rsplit("/", 1)[-1][:200]
    known = {r["seller_article"] for r in db.query_all(
        "SELECT seller_article FROM catalog_marketplace WHERE marketplace = %s", (mp,))}
    mp_label = TPL_MP_NAMES[mp]
    errors, exists, added = [], [], []
    with db.transaction() as tx:
        _lock_mp(tx, mp)
        seen, todo = {}, []
        for row_no, text in t["rows"]:
            sa = util.canon_article(text)
            if sa in seen:
                errors.append({"row": row_no, "article": text, "reason": "duplicate",
                               "detail": f"этот артикул уже есть в строке {seen[sa]} — записан один раз"})
                continue
            seen[sa] = row_no
            if sa not in known:
                errors.append({"row": row_no, "article": text, "reason": "not_found",
                               "detail": f"артикула нет в «Товарах» {mp_label} — проверьте, нет ли ошибки в написании"})
                continue
            todo.append((row_no, text, sa))
        others = {}
        if todo:
            for r in tx.query_all(
                    "SELECT seller_article, name, date_from, date_to_plan, date_to_fact, "
                    "COALESCE(date_to_fact, date_to_plan) AS date_to FROM rnp_promos "
                    "WHERE marketplace = %s AND seller_article = ANY(%s) AND cancelled_at IS NULL "
                    "AND date_from <= %s AND COALESCE(date_to_fact, date_to_plan) >= %s ORDER BY date_from",
                    (mp, [sa for _, _, sa in todo], d_to, d_from)):
                others.setdefault(r["seller_article"], []).append(r)
        new = []
        for row_no, text, sa in todo:
            hit = others.get(sa)
            if not hit:
                new.append(sa)
                added.append(text)
            elif any(o["name"] == name and o["date_from"] == d_from and o["date_to_plan"] == d_to
                     and not o["date_to_fact"] for o in hit):
                exists.append({"row": row_no, "article": text})
            else:
                o = hit[0]
                errors.append({"row": row_no, "article": text, "reason": "overlap",
                               "detail": f"«{o['name']}» {o['date_from']:%d.%m.%Y} — {o['date_to']:%d.%m.%Y}: "
                                         "даты пересекаются, две акции одновременно быть не может"})
        if new:
            ids = [r["id"] for r in tx.query_all(
                "INSERT INTO rnp_promos (marketplace, seller_article, name, date_from, date_to_plan, created_by) "
                "SELECT %s, a, %s, %s, %s, %s FROM unnest(%s::text[]) AS a RETURNING id",
                (mp, name, d_from, d_to, user["id"], new))]
            changes = json.dumps({"name": name, "date_from": d_from, "date_to_plan": d_to, "file": fname or "файл"},
                                 ensure_ascii=False, default=str)
            tx.execute("INSERT INTO rnp_promo_events (promo_id, action, changes, user_id) "
                       "SELECT i, 'create', %s::jsonb, %s FROM unnest(%s::bigint[]) AS i",
                       (changes, user["id"], ids))
    errors.sort(key=lambda e: e["row"])
    return {"ok": True, "file": fname, "marketplace": mp, "name": name,
            "date_from": d_from.isoformat(), "date_to_plan": d_to.isoformat(), "days": (d_to - d_from).days + 1,
            "rows": len(t["rows"]), "added": added, "exists": exists, "errors": errors}
