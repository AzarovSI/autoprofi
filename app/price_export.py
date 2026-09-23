"""Read-only price-index workbook, styled like the summary sales report."""
import datetime as dt
import io
import math
import xml.etree.ElementTree as ET
from zoneinfo import ZoneInfo
from zipfile import ZipFile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import util

DASH = "–"
REPORT_TITLE = "Индекс цен на маркетплейсах"


def export_time():
    return dt.datetime.now(ZoneInfo("Europe/Moscow"))


def text_cell(cell, value):
    """Never interpret catalog text as a spreadsheet formula."""
    cell.value = str(value or "")
    cell.data_type = "s"


def _number(value):
    if value is None:
        return None
    value = float(value)
    return value if math.isfinite(value) else None


def _save_with_caches(wb, caches):
    """Keep auditable formulas AND correct values for non-calculating viewers.

    Only our generated sheet/formula cells are patched; no input XLSX is parsed.
    Excel still recalculates formulas when the workbook is opened or edited.
    """
    raw = io.BytesIO()
    wb.save(raw)
    result = io.BytesIO()
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(raw) as src, ZipFile(result, "w") as dest:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(ns + "c"):
                    ref = cell.get("r")
                    if ref not in caches:
                        continue
                    value = caches[ref]
                    node = cell.find(ns + "v")
                    if node is None:
                        node = ET.SubElement(cell, ns + "v")
                    if value is None:
                        cell.set("t", "str")
                        node.text = DASH
                    else:
                        cell.attrib.pop("t", None)
                        node.text = repr(value)
                data = ET.tostring(root, encoding="utf-8")
            dest.writestr(info, data)
    result.seek(0)
    return result


def index_workbook(items, period=None, show_cost=False, search="", generated_at=None):
    """The same source snapshot/formulas as the on-screen Price Index."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Индекс цен"
    columns = [("article", "Иерархия / Артикул", 46)]
    if show_cost:
        columns.append(("cost", "Себестоимость, ₽", 18))
    for prefix in ("wb", "oz"):
        if prefix == "oz":
            columns.append(("buyer_pi", "Pi покупатель\nOzon / WB", 16))
        columns.extend([
            (prefix + "_base", "Базовая, ₽", 16),
            (prefix + "_pi", "Pi", 11),
            (prefix + "_upload_price", "Загружаемая, ₽", 18),
            (prefix + "_spp_pct", "СПП / соинвест, %", 16),
            (prefix + "_buyer_price", "Для покупателя, ₽", 18),
        ])
    pos = {key: i for i, (key, _, _) in enumerate(columns, 1)}
    last = len(columns)
    thin = Side(style="thin", color="BFC9D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    caches = {}

    def band(row, value, fill, color, size=10):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        c = ws.cell(row, 1)
        text_cell(c, value)
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(name="Calibri", size=size, bold=row == 1, color=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    generated_at = generated_at or export_time()
    band(1, f"{REPORT_TITLE} {generated_at:%d.%m.%Y}", "1F3B57", "FFFFFF", 15)
    period_label = (f"{period['date_from']} — {period['date_to']}" if period else "Весь период")
    formed = generated_at.strftime("%d.%m.%Y %H:%M МСК")
    band(2, f"Период: {period_label}. Срез на конец периода. Товаров: {len(items)}. Сформировано: {formed}.",
         "FFFFFF", "555555")
    band(3, "Pi = загружаемая / базовая; для АВТОПРОФИ: покупатель / базовая. "
         "Pi покупатель = Ozon / WB. Прочерк: нет данных. ≈: оценочная СПП."
         + (f" Поиск: {search}" if search else ""), "FFFFFF", "555555")
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 30

    for col, (key, label, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        for row in (4, 5):
            c = ws.cell(row, col)
            c.fill = PatternFill("solid", fgColor="2E5A87")
            c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c.alignment = center
            c.border = border
        if key in ("article", "cost", "buyer_pi"):
            ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
            text_cell(ws.cell(4, col), label)
        else:
            text_cell(ws.cell(5, col), label)
    for prefix, label, fill in (("wb", "Wildberries", "DCE6F1"), ("oz", "Ozon", "EAF1F8")):
        start, end = pos[prefix + "_base"], pos[prefix + "_buyer_price"]
        ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
        for col in range(start, end + 1):
            c = ws.cell(4, col)
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(name="Calibri", size=10, bold=True, color="1F3B57")
        text_cell(ws.cell(4, start), label)
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 32

    root = {"children": {}, "items": []}
    for item in items:
        node = root
        path = [item.get("l1") or "Не распределены по группам"]
        if item.get("l1") and item.get("l2"):
            path.append(item["l2"])
            if item.get("l3"):
                path.append(item["l3"])
        for name in path:
            node = node["children"].setdefault(name, {"children": {}, "items": []})
        node["items"].append(item)

    def ratio(row, key, numerator_key, denominator_key, threshold=None, blue=False):
        cell = ws.cell(row, pos[key])
        num = ws.cell(row, pos[numerator_key])
        den = ws.cell(row, pos[denominator_key])
        n, d = num.coordinate, den.coordinate
        cell.value = f'=IF(AND(ISNUMBER({n}),ISNUMBER({d}),{d}>0),{n}/{d},"{DASH}")'
        nv, dv = num.value, den.value
        caches[cell.coordinate] = nv / dv if isinstance(nv, (int, float)) and isinstance(dv, (int, float)) and dv > 0 else None
        cell.number_format = "0.00"
        ref = cell.coordinate
        if threshold is not None:
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=[f"AND(ISNUMBER({ref}),ROUND({ref},2)<{threshold})"],
                font=Font(color="9C0006"), stopIfTrue=True))
        if blue:
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=[f"AND(ISNUMBER({ref}),ROUND({ref},2)>1.1)"],
                font=Font(color="0563C1")))

    def write_node(node, depth=0):
        names = sorted(node["children"], key=lambda name: (
            (name == "Не распределены по группам", util.l1_sort_key(name))
            if depth == 0 else util.cat_sort_key(name)))
        for name in names:
            row = ws.max_row + 1
            for col in range(1, last + 1):
                cell = ws.cell(row, col)
                cell.fill = PatternFill("solid", fgColor=("BDD7EE", "DDEBF7", "F2F7FC")[depth])
                cell.font = Font(name="Calibri", size=10, bold=True, color="1F3B57")
                cell.border = border
            text_cell(ws.cell(row, 1), name)
            ws.cell(row, 1).alignment = Alignment(indent=depth, vertical="center")
            ws.row_dimensions[row].outlineLevel = depth
            ws.row_dimensions[row].height = 22
            write_node(node["children"][name], depth + 1)
        for item in sorted(node["items"], key=lambda it: util.cat_sort_key(it.get("seller_article"))):
            row = ws.max_row + 1
            for col in range(1, last + 1):
                c = ws.cell(row, col)
                c.font = Font(name="Calibri", size=10, color="333333")
                c.border = border
                c.alignment = right
                c.number_format = "#,##0"
            text_cell(ws.cell(row, 1), item.get("seller_article"))
            ws.cell(row, 1).alignment = Alignment(indent=depth, vertical="center")
            ws.row_dimensions[row].outlineLevel = depth
            ws.row_dimensions[row].height = 21
            if show_cost:
                c = ws.cell(row, pos["cost"])
                c.value = _number(item.get("cost")) if item.get("cost") is not None else DASH
                c.number_format = "#,##0.00"
            avto = "АВТОПРОФИ" in (item.get("l1") or "").upper()
            for prefix, base in (("wb", "price_wb"), ("oz", "price_ozon")):
                for suffix in ("base", "upload_price", "spp_pct", "buyer_price"):
                    key = prefix + "_" + suffix
                    value = _number(item.get(base if suffix == "base" else key))
                    c = ws.cell(row, pos[key])
                    c.value = value if value is not None else DASH
                    if suffix == "spp_pct":
                        estimated = item.get(prefix + "_spp_is_estimated")
                        c.number_format = '"≈"0.0%' if estimated else "0.0%"
                        if estimated:
                            c.comment = Comment("СПП оценочная: из ближайшего известного дня.", "ТД АВТОПРОФИ")
                ratio(row, prefix + "_pi", prefix + ("_buyer_price" if avto else "_upload_price"),
                      prefix + "_base", 1.15 if avto else 1, not avto)
            ratio(row, "buyer_pi", "oz_buyer_price", "wb_buyer_price", 1.03)

    write_node(root)
    if not items:
        text_cell(ws.cell(6, 1), "Ничего не найдено")
    ws.freeze_panes = "B6"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.print_title_rows = "1:5"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(last)}{ws.max_row}"
    return _save_with_caches(wb, caches)
