# -*- coding: utf-8 -*-
"""Загрузчик отчёта МОНИТОРИНГА ЦЕН КОНКУРЕНТОВ (PriceVA) для WILDBERRIES.

Полный аналог ozon_priceva_loader, но пишет в wb_daily_sales. Отчёт PriceVA
(«priceva_report_full_*.xlsx») ОДИН И ТОТ ЖЕ по формату для Ozon и WB, но
ДАННЫЕ В НЁМ РАЗНЫЕ (цены конкурентов на разных площадках). Эти данные
наполняют две метрики товара в разделе «РНП заказы» (WB):
    • «Средняя цена конкурентов»    → wb_daily_sales.comp_price_avg
    • «Минимальная цена конкурентов» → wb_daily_sales.comp_price_min

Формат файла (выгрузка PriceVA):
    • лист «Отчёт» (берётся первый лист, если имя другое);
    • строка 1 — заголовки (колонки «Мин./Средняя цена» имеют объединённые
      ячейки — читаем по ИМЕНИ заголовка, не по позиции);
    • данные со 2-й строки. Ключ строки — «Артикул».
    • ДАТЫ В ФАЙЛЕ НЕТ — она передаётся отдельным параметром при загрузке
      (за какую дату мониторинга загружаются цены). Это ОТЛИЧИЕ от отчёта
      «Воронка» WB (там дата в файле) — здесь дату задаём вручную, как у Ozon.

Что берём из отчёта (поиск заголовка по имени, с запасными вариантами):
    • «Артикул»            → seller_article (сопоставление со справочником);
    • «Мин. цена руб.»     → comp_price_min;
    • «Средняя цена руб.»  → comp_price_avg.

Сопоставление артикулов — через canon_article (util.py): trim + схлоп
внутренних пробелов + гомоглифы + UPPER. Разные пробелы/регистр не мешают.
Пишем цены ТОЛЬКО по товарам, которые есть в общем справочнике дашборда
(catalog_items). Товары не из справочника — пропускаем, их число отдаём.

Запись в wb_daily_sales за выбранную дату:
    • если строка продаж (date, seller_article) уже есть — UPDATE цен;
    • если строки нет — создаём строку-ЗАГЛУШКУ (заказы пустые, только
      цены конкурентов), чтобы цена сохранилась в любом случае. item_name
      берём из справочника (catalog_items.sample_name).
UPSERT по уникальному ключу (date, seller_article) ON CONFLICT.

Каждая загрузка фиксируется в report_uploads (marketplace='Wildberries',
period_kind='competitors').
"""
import os
import io
import zipfile
import datetime

import openpyxl

from . import util

MP_NAME = "Wildberries"
SHEET_NAME = "Отчёт"
HEADER_ROW = 1  # строка заголовков

# Заголовки колонок — проверка формата ПО ИМЕНИ (не по позиции).
# Для каждой логической колонки — список допустимых названий (первый
# найденный выигрывает). Это защищает от мелких расхождений в выгрузке.
ART_HEADERS = ["Артикул"]
MIN_HEADERS = ["Мин. цена руб.", "Минимальная цена руб.", "Мин. цена руб",
               "Минимальная цена руб"]
AVG_HEADERS = ["Средняя цена руб.", "Средняя цена в руб.", "Средняя цена руб",
               "Средняя цена в руб"]


def _to_num(v):
    """Число из ячейки; пусто/«-»/нечисло → None."""
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _build_colidx(ws, header_row):
    """Словарь {название заголовка -> номер колонки}. Первое вхождение."""
    idx = {}
    max_col = ws.max_column or 20
    for c in range(1, max_col + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip():
            idx.setdefault(str(v).strip(), c)
    return idx


def _find_col(idx, candidates):
    """Найти номер колонки по списку допустимых названий."""
    for name in candidates:
        if name in idx:
            return idx[name]
    return None


def _journal(conn, path, *, status, message="", rows_loaded=0,
             period_start=None, period_end=None, year=None,
             period_text="", upload_id=None):
    """Создать/обновить запись в report_uploads (period_kind='competitors')."""
    try:
        with conn.cursor() as cur:
            if upload_id is None:
                cur.execute(
                    "INSERT INTO report_uploads (marketplace, year, period_text, "
                    "period_start, period_end, source_file, rows_loaded, status, "
                    "message, period_kind) "
                    "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id",
                    (MP_NAME, year, period_text, period_start, period_end,
                     os.path.basename(path), rows_loaded, status,
                     (message or "")[:500], "competitors"))
                upload_id = cur.fetchone()[0]
            else:
                cur.execute(
                    "UPDATE report_uploads SET rows_loaded=%s, status=%s, "
                    "message=%s WHERE id=%s",
                    (rows_loaded, status, (message or "")[:500], upload_id))
        conn.commit()
        return upload_id
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
        return upload_id


def _open_workbook(path):
    """Открыть xlsx; при необходимости починить zip (регистр служебных имён)."""
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except KeyError:
        pass
    canon = {
        "xl/sharedstrings.xml": "xl/sharedStrings.xml",
        "xl/workbook.xml": "xl/workbook.xml",
        "xl/styles.xml": "xl/styles.xml",
    }
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        lower_map = {n.lower(): n for n in names}
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            written = set()
            for n in names:
                zout.writestr(n, zin.read(n))
                written.add(n)
            for low, canon_name in canon.items():
                if canon_name in written:
                    continue
                src = lower_map.get(low)
                if src is not None:
                    zout.writestr(canon_name, zin.read(src))
    buf.seek(0)
    return openpyxl.load_workbook(buf, data_only=True)


def _parse_date_arg(date_str):
    """Разобрать дату мониторинга, переданную при загрузке.

    Принимает YYYY-MM-DD (из <input type=date>) или ДД.ММ.ГГГГ.
    """
    if not date_str:
        return None
    s = str(date_str).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def load_wb_priceva(conn, path, report_date=None):
    """Загрузить отчёт цен конкурентов PriceVA (WB) за указанную дату.

    conn — psycopg-соединение (без автокоммита).
    report_date — дата мониторинга (str YYYY-MM-DD/ДД.ММ.ГГГГ или date);
    обязательна, т.к. в файле даты нет.

    Возвращает dict {ok, rows_loaded, updated, inserted, skipped_not_in_catalog,
    report_date, ...}. Пишет comp_price_avg / comp_price_min в wb_daily_sales
    за эту дату (UPDATE существующих строк + INSERT строк-заглушек).
    """
    if isinstance(report_date, datetime.date):
        d = report_date
    else:
        d = _parse_date_arg(report_date)
    if d is None:
        msg = ("Не указана дата мониторинга. Выберите дату, за которую "
               "загружаются цены конкурентов.")
        _journal(conn, path, status="ОШИБКА", message=msg)
        return {"ok": False, "error": msg}

    wb = _open_workbook(path)
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb[wb.sheetnames[0]]

    idx = _build_colidx(ws, header_row=HEADER_ROW)
    c_art = _find_col(idx, ART_HEADERS)
    c_min = _find_col(idx, MIN_HEADERS)
    c_avg = _find_col(idx, AVG_HEADERS)
    missing = []
    if c_art is None:
        missing.append("Артикул")
    if c_min is None:
        missing.append("Мин. цена руб.")
    if c_avg is None:
        missing.append("Средняя цена руб.")
    if missing:
        msg = ("ОШИБКА ФОРМАТА: не найдены колонки в строке 1: "
               + ", ".join(missing))
        _journal(conn, path, status="ОШИБКА ФОРМАТА", message=msg)
        return {"ok": False, "error": msg}

    # --- сбор canon(article) -> (min, avg); последнее значение при дублях ---
    data = {}  # art_canon -> (comp_min, comp_avg)
    skipped_empty = 0
    max_row = ws.max_row or HEADER_ROW
    for r in range(HEADER_ROW + 1, max_row + 1):
        art_raw = ws.cell(r, c_art).value
        if art_raw is None or not str(art_raw).strip():
            skipped_empty += 1
            continue
        art_c = util.canon_article(art_raw)
        if not art_c:
            skipped_empty += 1
            continue
        cmin = _to_num(ws.cell(r, c_min).value)
        cavg = _to_num(ws.cell(r, c_avg).value)
        data[art_c] = (cmin, cavg)

    if not data:
        msg = "В файле нет ни одной валидной строки данных (пустой отчёт)."
        _journal(conn, path, status="ОШИБКА ФОРМАТА", message=msg)
        return {"ok": False, "error": msg}

    period_text = d.strftime('%d.%m.%Y')
    year = d.year

    upload_id = _journal(conn, path, status="В ПРОЦЕССЕ", period_text=period_text,
                         period_start=d, period_end=d, year=year)

    try:
        with conn.cursor() as cur:
            # 0) Фильтр по ОБЩЕМУ СПРАВОЧНИКУ (catalog_items). Артикулы в
            #    справочнике уже канонические, но UPPER() на всякий случай.
            #    Заодно берём sample_name — для item_name строк-заглушек.
            cur.execute(
                "SELECT upper(seller_article), seller_article, "
                "COALESCE(sample_name,'') FROM catalog_items")
            catalog = {}  # art_canon -> (art_orig, name)
            for au, orig, name in cur.fetchall():
                catalog[au] = (orig, name)

            filtered = {ac: v for ac, v in data.items() if ac in catalog}
            skipped_not_in_catalog = len(data) - len(filtered)

            if not filtered:
                msg = (f"Дата {period_text}; в файле: {len(data)}; "
                       f"ни один артикул не найден в справочнике дашборда.")
                cur.execute(
                    "UPDATE report_uploads SET rows_loaded=0, status='OK', "
                    "message=%s WHERE id=%s", (msg[:500], upload_id))
                conn.commit()
                return {
                    "ok": True, "marketplace": MP_NAME, "rows_loaded": 0,
                    "updated": 0, "inserted": 0,
                    "skipped_not_in_catalog": skipped_not_in_catalog,
                    "file_total": len(data), "report_date": d.isoformat(),
                    "dates": [d.isoformat()], "period_text": period_text,
                    "skipped_empty": skipped_empty, "upload_id": upload_id,
                }

            # 1) UPSERT цен за дату. Ключ уникальности (date, seller_article).
            #    Пишем оригинальный (канонический) артикул из справочника —
            #    так строки-заглушки корректно джойнятся со справочником.
            #    ON CONFLICT: обновляем ТОЛЬКО цены конкурентов, остальные
            #    поля существующих строк продаж не трогаем.
            rows = []
            for ac, (cmin, cavg) in filtered.items():
                art_orig, name = catalog[ac]
                rows.append((d, art_orig, name or None, cavg, cmin))

            # Считаем, сколько строк уже существует за дату (для отчёта
            # updated/inserted) — до UPSERT.
            arts_orig = [r[1] for r in rows]
            cur.execute(
                "SELECT count(*) FROM wb_daily_sales "
                "WHERE date=%s AND seller_article = ANY(%s)",
                (d, arts_orig))
            existing_before = cur.fetchone()[0]

            cur.executemany(
                "INSERT INTO wb_daily_sales "
                "(date, seller_article, item_name, comp_price_avg, comp_price_min) "
                "VALUES (%s,%s,%s,%s,%s) "
                "ON CONFLICT (date, seller_article) DO UPDATE SET "
                "comp_price_avg = EXCLUDED.comp_price_avg, "
                "comp_price_min = EXCLUDED.comp_price_min",
                rows)

            total = len(rows)
            updated = existing_before
            inserted = total - existing_before

            msg = (f"Дата {period_text}; в файле: {len(data)}; "
                   f"записано цен: {total} (обновлено строк продаж: {updated}, "
                   f"создано строк-заглушек: {inserted})")
            if skipped_not_in_catalog:
                msg += f"; вне справочника (пропущено): {skipped_not_in_catalog}"
            if skipped_empty:
                msg += f"; пустых: {skipped_empty}"
            cur.execute(
                "UPDATE report_uploads SET rows_loaded=%s, status='OK', "
                "message=%s WHERE id=%s", (total, msg[:500], upload_id))
        conn.commit()
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        user_msg = "Не удалось сохранить отчёт цен конкурентов в базу."
        _journal(conn, path, status="ОШИБКА", message=str(exc),
                 upload_id=upload_id)
        return {"ok": False, "error": user_msg, "upload_id": upload_id}

    return {
        "ok": True,
        "marketplace": MP_NAME,
        "rows_loaded": total,
        "updated": updated,
        "inserted": inserted,
        "file_total": len(data),
        "skipped_not_in_catalog": skipped_not_in_catalog,
        "report_date": d.isoformat(),
        "dates": [d.isoformat()],
        "period_text": period_text,
        "skipped_empty": skipped_empty,
        "upload_id": upload_id,
    }
