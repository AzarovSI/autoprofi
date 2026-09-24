# -*- coding: utf-8 -*-
"""Загрузчик отчёта ОЗОН «Список товаров» (рейтинг + отзывы).

Источник данных для двух метрик матрицы «РНП продажи — OZON»:
    • «Рейтинг товара»       ← колонка «Рейтинг»  (формат '4.80 → 4.8)
    • «Количество отзывов»   ← колонка «Отзывы»   (целое)

В отчёте ДАТЫ НЕТ — она передаётся отдельным параметром при загрузке
(дата, на которую действителен отчёт), как в загрузке остатков.

Формат файла (Личный кабинет Озон → Товары → Список товаров →
скачать отчёт на дату):
    • XLSX (новый формат) ИЛИ CSV (старый) — определяется автоматически;
    • в XLSX 1-я строка — титул «Отчёт по товарам», заголовки во 2-й;
      в CSV заголовки в 1-й. Строка заголовков ищется ПО ИМЕНИ
      колонки «Артикул», поэтому оба формата читаются корректно;
    • ключевые колонки по ИМЕНИ: «Артикул», «Рейтинг», «Отзывы».

Данные пишутся ТОЛЬКО по товарам, которые есть в общем справочнике
дашборда (catalog_items). При отсутствии заказов создаётся строка за указанную
дату только с рейтингом и отзывами; показатели заказов остаются NULL.
Метрика рейтинга/отзывов — snapshot (последнее значение на дату).

Каждая загрузка фиксируется в report_uploads (marketplace='Ozon',
period_kind='reviews').
"""
import os
import io
import csv
import zipfile
import datetime

import openpyxl

MP_NAME = "Ozon"
PERIOD_KIND = "reviews"

# Ключевые заголовки — проверка формата ПО ИМЕНИ, не по позиции.
REQUIRED = ["Артикул", "Рейтинг", "Отзывы"]


def _parse_rating(v):
    """'4.80 → 4.8; пусто/нечисло → None."""
    if v is None:
        return None
    s = str(v).strip().lstrip("'").strip()
    if not s or s == "-":
        return None
    s = s.replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def _parse_int(v):
    """Целое число отзывов; пусто/нечисло → None."""
    if v is None:
        return None
    s = str(v).strip().lstrip("'").strip()
    if not s or s == "-":
        return None
    s = s.replace("\xa0", "").replace(" ", "").replace(",", "")
    try:
        return int(float(s))
    except ValueError:
        return None


def _parse_date_arg(date_str):
    """Дата из строки: YYYY-MM-DD | ДД.ММ.ГГГГ | ДД.ММ.ГГ → date | None."""
    if not date_str:
        return None
    s = str(date_str).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _journal(conn, path, *, status, message="", rows_loaded=0,
             period_start=None, period_end=None, year=None,
             period_text="", upload_id=None):
    """Создать/обновить запись в report_uploads (period_kind='reviews')."""
    try:
        with conn.cursor() as cur:
            if upload_id is None:
                cur.execute(
                    "INSERT INTO report_uploads (marketplace, year, period_text, "
                    "period_start, period_end, source_file, rows_loaded, status, "
                    "message, period_kind) "
                    "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id",
                    (MP_NAME, year, period_text, period_start, period_end,
                     os.path.basename(path), rows_loaded, status,
                     (message or "")[:500], PERIOD_KIND))
                upload_id = cur.fetchone()[0]
            else:
                cur.execute(
                    "UPDATE report_uploads SET rows_loaded=%s, status=%s, "
                    "message=%s WHERE id=%s",
                    (rows_loaded, status, (message or "")[:500], upload_id))
        conn.commit()
        return upload_id
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
        return upload_id


def _open_workbook(path):
    """Открыть xlsx; при необходимости починить zip выгрузки Ozon/1С.

    Выгрузка иногда содержит служебные файлы с неверным регистром имён
    (xl/SharedStrings.xml вместо xl/sharedStrings.xml), из-за чего
    openpyxl падает с KeyError. Пробуем открыть напрямую, при ошибке —
    пересобираем архив в памяти с именами в нужном регистре.
    """
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except KeyError:
        pass
    canon = {
        "xl/sharedstrings.xml": "xl/sharedStrings.xml",
        "xl/workbook.xml": "xl/workbook.xml",
        "xl/styles.xml": "xl/styles.xml",
    }
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        lower_map = {n.lower(): n for n in names}
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            written = set()
            for n in names:
                zout.writestr(n, zin.read(n))
                written.add(n)
            for low, canon_name in canon.items():
                if canon_name in written:
                    continue
                src = lower_map.get(low)
                if src is not None:
                    zout.writestr(canon_name, zin.read(src))
    buf.seek(0)
    return openpyxl.load_workbook(buf, data_only=True)


def _is_xlsx(path):
    """XLSX по расширению или по сигнатуре ZIP (PK\\x03\\x04)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return True
    if ext in (".csv", ".txt"):
        return False
    try:
        with open(path, "rb") as f:
            return f.read(4)[:2] == b"PK"
    except Exception:
        return False


def _read_xlsx_grid(path):
    """Все строки первого листа XLSX как список списков строк."""
    wb = _open_workbook(path)
    try:
        ws = wb[wb.sheetnames[0]]
        grid = []
        for row in ws.iter_rows(values_only=True):
            grid.append([("" if c is None else str(c)) for c in row])
        return grid
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _read_csv_grid(path):
    """Все строки CSV как список списков строк.

    UTF-8 (с BOM или без), разделитель «;» (запасной «,»).
    """
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        rd = csv.reader(f, delimiter=delim)
        return [row for row in rd if row]


def _find_header_row(grid):
    """Индекс строки заголовков — первой, где есть колонка «Артикул».

    В CSV заголовки в 1-й строке; в новом XLSX «Отчёт по товарам»
    занимает 1-ю строку, а заголовки — во 2-й. Ищем по имени, а не
    по позиции, поэтому оба формата работают без завязки на номер строки.
    """
    for i, row in enumerate(grid[:15]):
        names = {str(c).strip() for c in row if c is not None}
        if "Артикул" in names:
            return i
    return None


def _read_report_rows(path):
    """Прочитать отчёт (XLSX или CSV) → (header, rows).

    Возвращает заголовки и строки данных, автоматически определяя формат
    и пропуская титульные строки над таблицей.
    """
    grid = _read_xlsx_grid(path) if _is_xlsx(path) else _read_csv_grid(path)
    if not grid:
        return None, []
    h = _find_header_row(grid)
    if h is None:
        return None, []
    return grid[h], grid[h + 1:]


def load_ozon_reviews(conn, path, report_date=None):
    """Загрузить отчёт «Список товаров» Ozon (рейтинг + отзывы) за дату.

    conn — psycopg-соединение (без автокоммита).
    report_date — дата, на которую действителен отчёт (str или date);
    обязательна, т.к. в файле даты нет.

    Пишет rating и delivery_time_hours (=количество отзывов) в
    ozon_daily_sales независимо от наличия заказов за эту дату,
    только по товарам из справочника дашборда (catalog_items).
    """
    if isinstance(report_date, datetime.date):
        d = report_date
    else:
        d = _parse_date_arg(report_date)
    if d is None:
        msg = ("Не указана дата отчёта. Выберите дату, на которую "
               "действителен загружаемый отчёт.")
        _journal(conn, path, status="ОШИБКА", message=msg)
        return {"ok": False, "error": msg}

    header, rows = _read_report_rows(path)
    if not header:
        msg = "Файл пустой или не удалось найти строку заголовков (нужна колонка «Артикул»)."
        _journal(conn, path, status="ОШИБКА ФОРМАТА", message=msg)
        return {"ok": False, "error": msg}

    # индекс колонок по имени
    idx = {}
    for i, name in enumerate(header):
        nm = (name or "").strip()
        if nm:
            idx.setdefault(nm, i)
    missing = [name for name in REQUIRED if name not in idx]
    if missing:
        msg = ("ОШИБКА ФОРМАТА: не найдены колонки: " + ", ".join(missing))
        _journal(conn, path, status="ОШИБКА ФОРМАТА", message=msg)
        return {"ok": False, "error": msg}

    i_art = idx["Артикул"]
    i_rating = idx["Рейтинг"]
    i_reviews = idx["Отзывы"]

    # --- сбор артикул -> (rating, reviews). Последнее значение при дублях ---
    data = {}  # art_upper -> (art_original, rating, reviews)
    skipped_empty = 0
    for row in rows:
        if i_art >= len(row):
            skipped_empty += 1
            continue
        art_raw = row[i_art]
        if art_raw is None or not str(art_raw).strip():
            skipped_empty += 1
            continue
        art = str(art_raw).strip()
        rating = _parse_rating(row[i_rating]) if i_rating < len(row) else None
        reviews = _parse_int(row[i_reviews]) if i_reviews < len(row) else None
        data[art.upper()] = (art, rating, reviews)

    if not data:
        msg = "В файле нет ни одной валидной строки данных."
        _journal(conn, path, status="ОШИБКА ФОРМАТА", message=msg)
        return {"ok": False, "error": msg}

    period_text = d.strftime("%d.%m.%Y")
    year = d.year
    upload_id = _journal(conn, path, status="В ПРОЦЕССЕ", period_text=period_text,
                         period_start=d, period_end=d, year=year)

    try:
        with conn.cursor() as cur:
            # 0) Фильтр по ОБЩЕМУ СПРАВОЧНИКУ дашборда (catalog_items).
            cur.execute("SELECT upper(seller_article), seller_article FROM catalog_items")
            catalog = dict(cur.fetchall())
            filtered = {au: v for au, v in data.items() if au in catalog}
            skipped_not_in_catalog = len(data) - len(filtered)

            synced = 0
            if filtered:
                # Независимый snapshot: не ждём появления строки заказов.
                # Канонический артикул берём из справочника, не из файла.
                vals = []
                params = []
                for au, (art_orig, rating, reviews) in filtered.items():
                    vals.append("(%s,%s,%s,%s)")
                    params.extend([d, catalog[au], rating, reviews])
                sql = (
                    "INSERT INTO ozon_daily_sales "
                    "(date, seller_article, rating, delivery_time_hours) VALUES "
                    + ",".join(vals) +
                    " ON CONFLICT (date, seller_article) DO UPDATE SET "
                    "rating=EXCLUDED.rating, "
                    "delivery_time_hours=EXCLUDED.delivery_time_hours")
                cur.execute(sql, params)
                synced = cur.rowcount

            matched = len(filtered)
            msg = (f"Дата отчёта {period_text}; в файле: {len(data)}; "
                   f"в справочнике дашборда: {matched}; "
                   f"синхр. в РНП Озон (рейтинг+отзывы): {synced}")
            if skipped_not_in_catalog:
                msg += f"; вне справочника (пропущено): {skipped_not_in_catalog}"
            if skipped_empty:
                msg += f"; пустых: {skipped_empty}"
            cur.execute(
                "UPDATE report_uploads SET rows_loaded=%s, status='OK', "
                "message=%s WHERE id=%s",
                (synced, msg[:500], upload_id))
        conn.commit()
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        user_msg = "Не удалось сохранить отчёт (рейтинг/отзывы) в базу."
        _journal(conn, path, status="ОШИБКА", message=str(exc),
                 upload_id=upload_id)
        return {"ok": False, "error": user_msg, "upload_id": upload_id}

    return {
        "ok": True,
        "marketplace": MP_NAME,
        "rows_loaded": synced,
        "synced_ozon": synced,
        "matched": len(filtered),
        "file_total": len(data),
        "skipped_not_in_catalog": skipped_not_in_catalog,
        "report_date": d.isoformat(),
        "period_text": period_text,
        "skipped_empty": skipped_empty,
        "upload_id": upload_id,
    }
