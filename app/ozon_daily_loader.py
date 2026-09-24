# -*- coding: utf-8 -*-
"""Загрузчик ДНЕВНОГО API-файла Ozon (лист «Данные») в ozon_daily_sales.

Отдельно от loader_core, т.к. формат другой: дневные данные (одна строка =
день × артикул), 16 колонок, заголовки в строке 1, данные со 2-й строки.

Перезагрузка за уже загруженные даты = перезапись:
    DELETE FROM ozon_daily_sales WHERE date BETWEEN min_date AND max_date
    затем INSERT батчами по 100 (лимит параметров psycopg3).

Артикулы, которых нет в справочнике, — upsert в catalog_marketplace (Ozon) и
catalog_items, чтобы данные не терялись (останутся «нераспределёнными» до
привязки категорий/статуса в Справочнике — это нормально).

Каждая загрузка фиксируется в report_uploads (marketplace='Ozon',
period_kind='daily').
"""
import os
import datetime

import openpyxl

try:
    from .util import canon_article
except ImportError:  # запуск напрямую
    from util import canon_article

MP_NAME = "Ozon"
SHEET_NAME = "Данные"

# Ключевые заголовки файла (проверка формата ПО ИМЕНИ, а не по позиции).
OZ_DAILY_REQUIRED = [
    "Дата", "Артикул", "SKU", "Заказы, шт", "Сумма заказов, ₽",
    "Отменено, шт", "Индекс цены (Pi)", "Скидка Ozon (СПП/соинвест), ₽",
    "СПП/соинвест, %", "Переходы в карточку", "CR в корзину, %",
    "CR в заказ, %", "Ср. позиция", "Остаток, шт",
]

# Маппинг: поле БД -> заголовок колонки файла. TEXT/ID-поля обрабатываются
# отдельно (date, seller_article, sku_ozon, item_name).
#
# ВАЖНО: сюда входят ТОЛЬКО поля, которые реально приносит дневной API-файл.
# «Рейтинг (отзывы)» в файле присутствует как заголовок, но всегда пустой —
# рейтинг/отзывы приходят из отдельного отчёта «Список товаров»
# (ozon_reviews_loader → rating, delivery_time_hours), поэтому rating НЕ
# мапится здесь, чтобы дневной файл не затирал его пустым значением.
METRIC_MAP = [
    ("orders_qty",     "Заказы, шт"),
    ("orders_rub",     "Сумма заказов, ₽"),
    ("cancels_qty",    "Отменено, шт"),
    ("price_index_pi", "Индекс цены (Pi)"),
    ("spp_rub",        "Скидка Ozon (СПП/соинвест), ₽"),
    ("spp_pct",        "СПП/соинвест, %"),
    ("card_visits",    "Переходы в карточку"),
    ("cr_cart_pct",    "CR в корзину, %"),
    ("cr_order_pct",   "CR в заказ, %"),
    ("avg_position",   "Ср. позиция"),
    ("stock_ozon_qty", "Остаток, шт"),
]

# Поля ozon_daily_sales, которые НАПОЛНЯЮТ ДРУГИЕ загрузчики/расчёты, а не
# дневной API-файл. При перезаписи дат (DELETE+INSERT) их значения нужно
# СОХРАНИТЬ по ключу (date, seller_article), иначе повторная загрузка
# дневного файла обнулит уже загруженные:
#   • ads_expense_rub, ctr_pct           — отчёт рекламы (ozon_ads_loader)
#   • rating, delivery_time_hours        — «Список товаров» (ozon_reviews_loader)
#   • stock_ap_qty                       — остатки ТД АВТОПРОФИ (ozon_stock_loader)
#   • avg_upload_price, price_for_buyer,
#     comp_price_avg, comp_price_min,
#     drr_total_pct                      — прочие источники/расчёты
PRESERVE_COLS = [
    "rating", "delivery_time_hours", "ads_expense_rub", "ctr_pct",
    "stock_ap_qty", "avg_upload_price", "price_for_buyer",
    "comp_price_avg", "comp_price_min", "drr_total_pct",
]

# Порядок колонок INSERT в ozon_daily_sales (без id/created_at — дефолты БД).
INSERT_COLS = (
    ["date", "seller_article", "sku_ozon", "item_name"]
    + [m[0] for m in METRIC_MAP]
    + ["upload_id"]
)


def _to_num(v):
    """Число из ячейки; пусто/«-»/нечисло → None."""
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    s = s.replace("%", "")
    try:
        return float(s)
    except ValueError:
        return None


def _to_date(v):
    """Дата из ячейки (datetime/date/строка ISO/ДД.ММ.ГГГГ). Нераспознанное → None."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_bigint(v):
    """float→bigint (int(round(v))); пусто → None."""
    n = _to_num(v)
    if n is None:
        return None
    try:
        return int(round(n))
    except (ValueError, OverflowError):
        return None


def _build_colidx(ws, header_row=1):
    """Индекс заголовок->номер колонки по строке header_row."""
    idx = {}
    max_col = ws.max_column or 16
    for c in range(1, max_col + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip():
            idx.setdefault(str(v).strip(), c)
    return idx


def _journal_error(conn, path, msg, period_start=None, period_end=None, year=None):
    """Записать ОШИБКУ в report_uploads (отдельная запись)."""
    period_text = ""
    if period_start and period_end:
        period_text = (f"{period_start.strftime('%d.%m.%Y')} - "
                       f"{period_end.strftime('%d.%m.%Y')}")
    try:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO report_uploads (marketplace, year, period_text, "
                "period_start, period_end, source_file, rows_loaded, status, "
                "message, period_kind) "
                "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                (MP_NAME, year, period_text, period_start, period_end,
                 os.path.basename(path), 0, "ОШИБКА ФОРМАТА", msg[:500], "daily"))
        conn.commit()
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass


def load_ozon_daily(conn, path):
    """Загрузить дневной API-файл Ozon в ozon_daily_sales.

    conn — psycopg-соединение (без автокоммита; функция сама commit/rollback).
    Возвращает dict {ok, rows_loaded, dates, unmatched, deleted, ...}.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        _journal_error(conn, path,
                       f"В файле нет листа «{SHEET_NAME}». Проверьте формат выгрузки.")
        return {"ok": False,
                "error": f"В файле нет листа «{SHEET_NAME}». Проверьте формат выгрузки."}

    idx = _build_colidx(ws, header_row=1)
    missing = [name for name in OZ_DAILY_REQUIRED if name not in idx]
    if missing:
        msg = "ОШИБКА ФОРМАТА: не найдены колонки: " + ", ".join(missing)
        _journal_error(conn, path, msg)
        return {"ok": False, "error": msg}

    c_date = idx["Дата"]
    c_art = idx["Артикул"]
    c_name = idx.get("Наименование")
    c_sku = idx.get("SKU")

    # --- чтение строк ---
    rows = []                # кортежи в порядке INSERT_COLS (без upload_id — добавим позже)
    dates = set()
    report_arts = {}         # canon_article -> item_name (для автодобавления в справочник)
    skipped_empty = 0
    skipped_baddate = 0

    max_row = ws.max_row or 1
    for r in range(2, max_row + 1):
        raw_date = ws.cell(r, c_date).value
        raw_art = ws.cell(r, c_art).value
        art_canon = canon_article(raw_art)
        d = _to_date(raw_date)

        # Пропускаем ПОЛНОСТЬЮ пустые строки (нет даты И нет артикула).
        if d is None and not art_canon:
            skipped_empty += 1
            continue
        # Строка с артикулом, но без распознанной даты — не можем ключевать.
        if d is None:
            skipped_baddate += 1
            continue
        if not art_canon:
            skipped_empty += 1
            continue

        name = ws.cell(r, c_name).value if c_name else None
        name_s = str(name).strip() if name is not None else None
        sku = _to_bigint(ws.cell(r, c_sku).value) if c_sku else None

        rec = {
            "date": d, "seller_article": art_canon,
            "sku_ozon": sku, "item_name": name_s,
        }
        for db_col, header in METRIC_MAP:
            c = idx.get(header)
            rec[db_col] = _to_num(ws.cell(r, c).value) if c else None

        dates.add(d)
        if art_canon not in report_arts:
            report_arts[art_canon] = name_s
        rows.append(rec)

    if not rows:
        msg = "В файле нет ни одной валидной строки данных (пустой отчёт)."
        _journal_error(conn, path, msg)
        return {"ok": False, "error": msg}

    min_date = min(dates)
    max_date = max(dates)
    period_text = (f"{min_date.strftime('%d.%m.%Y')} - "
                   f"{max_date.strftime('%d.%m.%Y')}")
    year = min_date.year

    # --- автодобавление артикулов в справочник ---
    # Существующие категории — чтобы понять, какие останутся нераспределёнными.
    existing_cat = {}
    with conn.cursor() as cur:
        cur.execute("SELECT seller_article, category_l1, category_l2, category_l3 "
                    "FROM catalog_items")
        for a, l1, l2, l3 in cur.fetchall():
            existing_cat[canon_article(a)] = (l1, l2, l3)

    new_arts = {a: nm for a, nm in report_arts.items() if a not in existing_cat}
    if new_arts:
        with conn.cursor() as cur:
            # батчевая вставка одним запросом (вместо N отдельных execute по сети)
            cur.executemany(
                "INSERT INTO catalog_items (seller_article, sample_name) "
                "VALUES (%s,%s) ON CONFLICT (seller_article) DO NOTHING",
                list(new_arts.items()))
        conn.commit()

    # Гарантируем строку в catalog_marketplace для Ozon по каждому артикулу (батчем).
    with conn.cursor() as cur:
        cur.executemany(
            "INSERT INTO catalog_marketplace (seller_article, marketplace) "
            "VALUES (%s,%s) ON CONFLICT (seller_article, marketplace) DO NOTHING",
            [(a, MP_NAME) for a in report_arts])
    conn.commit()

    # Нераспределённые (нет полной тройки категорий) — для отчёта пользователю.
    unmatched = sorted(
        a for a in report_arts
        if not all(existing_cat.get(a, (None, None, None)))
    )

    # --- журнал: создаём запись заранее (upload_id) ---
    with conn.cursor() as cur:
        cur.execute(
            "INSERT INTO report_uploads (marketplace, year, period_text, "
            "period_start, period_end, source_file, status, period_kind) "
            "VALUES (%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id",
            (MP_NAME, year, period_text, min_date, max_date,
             os.path.basename(path), "В ПРОЦЕССЕ", "daily"))
        upload_id = cur.fetchone()[0]

    # Кортежи для вставки в порядке INSERT_COLS.
    tuples = []
    for rec in rows:
        rec["upload_id"] = upload_id
        tuples.append(tuple(rec[col] for col in INSERT_COLS))

    # --- перезапись диапазона дат + вставка батчами ---
    # Перед DELETE сохраняем «чужие» поля (PRESERVE_COLS), которые наполняют
    # другие загрузчики (реклама/CTR, отзывы/рейтинг, остатки ТД и др.),
    # чтобы после INSERT восстановить их по ключу (date, seller_article).
    try:
        with conn.cursor() as cur:
            preserve_sql = ",".join(PRESERVE_COLS)
            cur.execute(
                f"SELECT date, seller_article, {preserve_sql} "
                "FROM ozon_daily_sales WHERE date BETWEEN %s AND %s",
                (min_date, max_date))
            preserved = {}  # (date, seller_article) -> {col: value}
            for row in cur.fetchall():
                d_key, art_key = row[0], row[1]
                vals = dict(zip(PRESERVE_COLS, row[2:]))
                # Сохраняем только если есть хотя бы одно непустое значение.
                if any(v is not None for v in vals.values()):
                    preserved[(d_key, art_key)] = vals

            cur.execute(
                "DELETE FROM ozon_daily_sales WHERE date BETWEEN %s AND %s",
                (min_date, max_date))
            deleted = cur.rowcount

            cols_sql = ",".join(INSERT_COLS)
            row_tmpl = "(" + ",".join(["%s"] * len(INSERT_COLS)) + ")"
            PAGE = 100
            for i in range(0, len(tuples), PAGE):
                batch = tuples[i:i + PAGE]
                values_sql = ",".join([row_tmpl] * len(batch))
                flat = [v for rr in batch for v in rr]
                cur.execute(
                    f"INSERT INTO ozon_daily_sales ({cols_sql}) VALUES {values_sql}",
                    flat)

            # Восстанавливаем сохранённые «чужие» поля для тех ключей, которые
            # снова есть в новой загрузке (те же date+seller_article).
            # ОДИН батчевый UPDATE ... FROM (VALUES ...) вместо N одиночных
            # UPDATE по сети (главное узкое место — было до ~1800 round-trip).
            restored = 0
            if preserved:
                new_keys = {(rec["date"], rec["seller_article"]) for rec in rows}
                upd_rows = [
                    ((d_key, art_key), vals)
                    for (d_key, art_key), vals in preserved.items()
                    if (d_key, art_key) in new_keys
                ]
                if upd_rows:
                    # Каждая VALUES-строка: date, seller_article, <PRESERVE_COLS...>.
                    # Типы задаём кастами в ПЕРВОЙ строке VALUES, чтобы PG знал
                    # типы столбцов (date, text, numeric...). Остальные — обычные %s.
                    n_pc = len(PRESERVE_COLS)
                    first_tmpl = ("(%s::date, %s::text, "
                                  + ",".join(["%s::numeric"] * n_pc) + ")")
                    rest_tmpl = "(" + ",".join(["%s"] * (2 + n_pc)) + ")"
                    vals_parts, flat = [], []
                    for i, ((d_key, art_key), vals) in enumerate(upd_rows):
                        vals_parts.append(first_tmpl if i == 0 else rest_tmpl)
                        flat.append(d_key)
                        flat.append(art_key)
                        flat.extend(vals[c] for c in PRESERVE_COLS)
                    set_sql = ",".join(f"{c}=v.{c}" for c in PRESERVE_COLS)
                    cols_list = "d, sa, " + ", ".join(PRESERVE_COLS)
                    cur.execute(
                        f"UPDATE ozon_daily_sales AS t SET {set_sql} "
                        f"FROM (VALUES {','.join(vals_parts)}) "
                        f"AS v({cols_list}) "
                        "WHERE t.date = v.d AND t.seller_article = v.sa",
                        flat)
                    restored = cur.rowcount

            msg = (f"Даты {period_text}; удалено старых: {deleted}; "
                   f"вставлено: {len(tuples)}")
            if restored:
                msg += f"; сохранено внешних полей: {restored}"
            if skipped_empty:
                msg += f"; пропущено пустых: {skipped_empty}"
            if skipped_baddate:
                msg += f"; без даты: {skipped_baddate}"
            if new_arts:
                msg += f"; новых артикулов: {len(new_arts)}"
            cur.execute(
                "UPDATE report_uploads SET rows_loaded=%s, status='OK', "
                "message=%s WHERE id=%s",
                (len(tuples), msg[:500], upload_id))
        conn.commit()
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        user_msg = ("Не удалось сохранить дневной отчёт Ozon в базу. "
                    "Проверьте файл на дубли по (дата, артикул).")
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "UPDATE report_uploads SET status='ОШИБКА', message=%s WHERE id=%s",
                    (str(exc)[:500], upload_id))
            conn.commit()
        except Exception:
            try:
                conn.rollback()
            except Exception:
                pass
        return {"ok": False, "error": user_msg, "upload_id": upload_id}

    return {
        "ok": True,
        "marketplace": MP_NAME,
        "rows_loaded": len(tuples),
        "deleted": deleted,
        "dates": [d.isoformat() for d in sorted(dates)],
        "period_text": period_text,
        "unmatched": unmatched,
        "new_articles": sorted(new_arts.keys()),
        "skipped_empty": skipped_empty,
        "skipped_baddate": skipped_baddate,
        "upload_id": upload_id,
    }
